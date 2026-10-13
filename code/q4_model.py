from __future__ import annotations

import argparse
import hashlib
import json
import math
import platform
import sys
import time as time_module
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]

import numpy as np
import pandas as pd
import scipy
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import csr_matrix, lil_matrix


N_DAYS = 365
N_INTERVALS = 144
RELEASE_HOURS = (0, 6, 12, 18)
RELEASE_STARTS = tuple(hour * 6 for hour in RELEASE_HOURS)
DELTA_HOURS = 1.0 / 6.0
ETA_CHARGE = 0.9
ETA_DISCHARGE = 0.9
SOC_INITIAL_KWH = 6000.0
SOC_MIN_KWH = 1200.0
SOC_MAX_KWH = 10800.0
INTERVAL_LIMIT_KWH = 5000.0 * DELTA_HOURS
EVALUATION_START_INDEX = 31
FINAL_DAY_INDEX = 364
SCENARIO_WINDOW_DAYS = 28
MIN_SCENARIO_DAYS = 7
GAMMA = 0.5
PRIMARY_TOLERANCE = 1e-6
PLANNING_CHECK_TOLERANCE = 2e-3
PHYSICAL_CHECK_TOLERANCE = 5e-5
EVENT_TOLERANCE = 1e-6
RANDOM_SEED = 2026
FORMAL_PRICE_INFORMATION = (
    "为完成题目给定数据下的离线重算，将附件4目标日完整价格路径视为规划输入；"
    "这不代表现实运行可提前获得未来实时价格。"
)

SOURCE_ROOT = PROJECT_ROOT / "data"
ACTUAL_PATH = SOURCE_ROOT / "附件2.xlsx"
FORECAST_PATH = SOURCE_ROOT / "附件3.xlsx"
FLOATING_PRICE_PATH = SOURCE_ROOT / "附件4.xlsx"


@dataclass(frozen=True)
class Q4Data:
    dates: list[date]
    load_kw: np.ndarray
    pv_kw: np.ndarray
    net_load_kwh: np.ndarray
    actual_price: np.ndarray
    official_hourly_kw: np.ndarray
    interval_starts: list[str]
    interval_ends: list[str]
    interval_labels: list[str]
    hashes: dict[str, str]


@dataclass(frozen=True)
class ForecastArchive:
    archive_id: str
    load_kw: np.ndarray
    price_hat: np.ndarray
    price_ratio: np.ndarray
    pv_kw: np.ndarray
    net_kwh: np.ndarray
    errors_kwh: np.ndarray
    valid: np.ndarray


@dataclass(frozen=True)
class ScenarioSet:
    net_kwh: np.ndarray
    price: np.ndarray
    probabilities: np.ndarray
    historical_day_indices: tuple[int, ...]
    degenerate: bool
    reason: str | None


@dataclass(frozen=True)
class PolicyConfig:
    strategy_id: str
    role: str
    question_variant: str
    forecast_source: str
    scenario_mode: str
    allow_adjustment: bool
    alpha: float | None = 0.8
    collect_evidence: bool = False


@dataclass
class LPSolution:
    contract: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    soc_end: np.ndarray
    emergency: np.ndarray
    unused: np.ndarray
    stage1_objective: float
    stage2_primary_objective: float
    throughput: float
    max_balance_residual: float
    max_storage_residual: float
    max_emergency_formula_residual: float
    max_adjustment_residual: float
    max_charge_discharge_overlap: float
    stage1_status: str
    stage2_status: str


@dataclass
class StrategyRun:
    config: PolicyConfig
    runtime_seconds: float
    metrics: dict[str, Any]
    validation: dict[str, Any]
    daily: pd.DataFrame
    ledger: pd.DataFrame | None
    adjustments: pd.DataFrame | None
    plan_calls: pd.DataFrame | None
    plan_trajectory: pd.DataFrame | None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def save_frame(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, index=False, encoding="utf-8-sig", float_format="%.10f")


def _clock(value: Any) -> str:
    if isinstance(value, time):
        return f"{value.hour:02d}:{value.minute:02d}"
    return str(value).strip()


def _fmt_minutes(minutes: int) -> str:
    if minutes == 1440:
        return "24:00"
    hour, minute = divmod(minutes, 60)
    return f"{hour:02d}:{minute:02d}"


def _expected_end_labels(index: int) -> set[str]:
    minutes = (index + 1) * 10
    if minutes == 1440:
        return {"0:00+1", "00:00+1", "24:00"}
    value = _fmt_minutes(minutes)
    return {value, value.lstrip("0")}


def _to_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()


def load_data() -> Q4Data:
    for path in (ACTUAL_PATH, FORECAST_PATH, FLOATING_PRICE_PATH):
        if not path.is_file():
            raise FileNotFoundError(path)

    actual_book = load_workbook(ACTUAL_PATH, read_only=True, data_only=True)
    if actual_book.sheetnames != ["小区负载", "光伏发电实际功率"]:
        raise ValueError(f"Unexpected attachment 2 sheets: {actual_book.sheetnames}")
    matrices: dict[str, np.ndarray] = {}
    dates: list[date] | None = None
    labels: list[str] | None = None
    for sheet_name in actual_book.sheetnames:
        rows = list(actual_book[sheet_name].values)
        if len(rows) != N_DAYS + 1 or len(rows[0]) != N_INTERVALS + 1:
            raise ValueError(f"Unexpected attachment 2 shape in {sheet_name}")
        current_dates = [_to_date(row[0]) for row in rows[1:]]
        current_labels = [_clock(value) for value in rows[0][1:]]
        if dates is None:
            dates, labels = current_dates, current_labels
        elif dates != current_dates or labels != current_labels:
            raise ValueError("Attachment 2 axes differ between sheets")
        matrices[sheet_name] = np.asarray([row[1:] for row in rows[1:]], dtype=float)
    assert dates is not None and labels is not None
    expected_dates = [date(2025, 1, 1) + timedelta(days=index) for index in range(N_DAYS)]
    if dates != expected_dates:
        raise ValueError("Attachment 2 does not cover every day of 2025 in order")
    for index, label in enumerate(labels):
        if label not in _expected_end_labels(index):
            raise ValueError(f"Attachment 2 time mismatch at interval {index + 1}")
    load_kw = matrices["小区负载"]
    pv_kw = matrices["光伏发电实际功率"]
    if not np.isfinite(load_kw).all() or not np.isfinite(pv_kw).all():
        raise ValueError("Attachment 2 contains missing/non-finite values")
    if np.any(load_kw < 0) or np.any(pv_kw < 0):
        raise ValueError("Attachment 2 physical powers must be nonnegative")

    price_book = load_workbook(FLOATING_PRICE_PATH, read_only=True, data_only=True)
    price_rows = list(price_book.active.values)
    if len(price_rows) != N_DAYS + 1 or len(price_rows[0]) != N_INTERVALS + 1:
        raise ValueError("Unexpected attachment 4 shape")
    price_dates = [_to_date(row[0]) for row in price_rows[1:]]
    price_labels = [_clock(value) for value in price_rows[0][1:]]
    if price_dates != dates or price_labels != labels:
        raise ValueError("Attachment 4 axes do not match attachment 2")
    actual_price = np.asarray([row[1:] for row in price_rows[1:]], dtype=float)
    if not np.isfinite(actual_price).all() or np.any(actual_price <= 0):
        raise ValueError("Attachment 4 prices must be finite and strictly positive")

    forecast_book = load_workbook(FORECAST_PATH, read_only=True, data_only=True)
    forecast_rows = list(forecast_book.active.values)
    expected_header = ("日期", "预报时刻", *[f"预报{lead}小时" for lead in range(1, 25)])
    if tuple(forecast_rows[0]) != expected_header or len(forecast_rows) != N_DAYS * 4 + 1:
        raise ValueError("Unexpected attachment 3 header or row count")
    official = np.full((N_DAYS, 4, 24), np.nan, dtype=float)
    current_date: date | None = None
    seen: set[tuple[date, int]] = set()
    release_labels = [f"{hour}:00" for hour in RELEASE_HOURS]
    for row in forecast_rows[1:]:
        if row[0] not in (None, ""):
            current_date = datetime.strptime(str(row[0]).strip(), "%Y-%m-%d").date()
        if current_date is None:
            raise ValueError("Attachment 3 begins with a blank date")
        try:
            release_index = release_labels.index(str(row[1]).strip())
        except ValueError as exc:
            raise ValueError(f"Unexpected release time {row[1]!r}") from exc
        day_index = (current_date - date(2025, 1, 1)).days
        if not 0 <= day_index < N_DAYS or (current_date, release_index) in seen:
            raise ValueError("Attachment 3 has duplicate or out-of-range release")
        values = np.asarray(row[2:], dtype=float)
        if len(values) != 24 or not np.isfinite(values).all() or np.any(values < 0):
            raise ValueError(f"Invalid attachment 3 forecast at {current_date} {row[1]}")
        official[day_index, release_index] = values
        seen.add((current_date, release_index))
    if np.isnan(official).any() or len(seen) != N_DAYS * 4:
        raise ValueError("Attachment 3 release coverage is incomplete")

    starts = [_fmt_minutes(index * 10) for index in range(N_INTERVALS)]
    ends = [_fmt_minutes((index + 1) * 10) for index in range(N_INTERVALS)]
    return Q4Data(
        dates=dates,
        load_kw=load_kw,
        pv_kw=pv_kw,
        net_load_kwh=(load_kw - pv_kw) * DELTA_HOURS,
        actual_price=actual_price,
        official_hourly_kw=official,
        interval_starts=starts,
        interval_ends=ends,
        interval_labels=[f"{start}-{end}" for start, end in zip(starts, ends)],
        hashes={
            "attachment2": sha256(ACTUAL_PATH),
            "attachment3": sha256(FORECAST_PATH),
            "attachment4": sha256(FLOATING_PRICE_PATH),
        },
    )


def _historical_indices(dates: list[date], day_index: int) -> np.ndarray:
    candidates = np.arange(max(0, day_index - 28), day_index, dtype=int)
    target_weekday = dates[day_index].weekday()
    same_weekday = np.asarray([idx for idx in candidates if dates[idx].weekday() == target_weekday], dtype=int)
    if len(same_weekday):
        return same_weekday
    target_weekend = target_weekday >= 5
    same_type = np.asarray(
        [idx for idx in candidates if (dates[idx].weekday() >= 5) == target_weekend], dtype=int
    )
    if len(same_type):
        return same_type
    return np.arange(max(0, day_index - 7), day_index, dtype=int)


def build_load_and_price_forecasts(data: Q4Data) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    load_hat = np.full_like(data.load_kw, np.nan)
    price_hat = np.full_like(data.actual_price, np.nan)
    valid_days = np.zeros(N_DAYS, dtype=bool)
    for day_index in range(1, N_DAYS):
        indices = _historical_indices(data.dates, day_index)
        if not len(indices):
            continue
        weights = np.exp(-(day_index - indices) / 7.0)
        load_hat[day_index] = np.average(data.load_kw[indices], axis=0, weights=weights)
        price_hat[day_index] = np.average(data.actual_price[indices], axis=0, weights=weights)
        valid_days[day_index] = True
    if np.any(price_hat[valid_days] <= 0):
        raise ValueError("Causal price forecasts must remain strictly positive")
    return load_hat, price_hat, valid_days


def _interpolate_official_day(data: Q4Data, day_index: int, release_index: int) -> np.ndarray:
    release_hour = RELEASE_HOURS[release_index]
    start = RELEASE_STARTS[release_index]
    if release_hour == 0:
        anchor = 0.0 if day_index == 0 else float(data.pv_kw[day_index - 1, -1])
    else:
        anchor = float(data.pv_kw[day_index, start - 1])
    knots = np.r_[anchor, data.official_hourly_kw[day_index, release_index]]
    targets = (np.arange(start, N_INTERVALS, dtype=float) + 1.0) * DELTA_HOURS - release_hour
    lower = np.minimum(np.floor(targets + 1e-12).astype(int), 23)
    fraction = targets - lower
    values = (1.0 - fraction) * knots[lower] + fraction * knots[lower + 1]
    result = np.full(N_INTERVALS, np.nan, dtype=float)
    result[start:] = np.maximum(values, 0.0)
    return result


def build_forecast_archive(data: Q4Data, source: str) -> ForecastArchive:
    load_hat, price_hat, valid_days = build_load_and_price_forecasts(data)
    pv_hat = np.full((N_DAYS, 4, N_INTERVALS), np.nan, dtype=float)
    if source == "official":
        for day_index in range(N_DAYS):
            for release_index in range(4):
                pv_hat[day_index, release_index] = _interpolate_official_day(data, day_index, release_index)
    elif source == "historical":
        for day_index in range(1, N_DAYS):
            indices = np.arange(max(0, day_index - 7), day_index, dtype=int)
            weights = np.exp(-(day_index - indices) / 7.0)
            value = np.maximum(0.0, np.average(data.pv_kw[indices], axis=0, weights=weights))
            for release_index, start in enumerate(RELEASE_STARTS):
                pv_hat[day_index, release_index, start:] = value[start:]
    else:
        raise ValueError(f"Unknown forecast source {source!r}")

    net_hat = np.full_like(pv_hat, np.nan)
    errors = np.full_like(pv_hat, np.nan)
    valid = np.zeros_like(pv_hat, dtype=bool)
    for day_index in range(1, N_DAYS):
        for release_index, start in enumerate(RELEASE_STARTS):
            net_hat[day_index, release_index, start:] = (
                load_hat[day_index, start:] - pv_hat[day_index, release_index, start:]
            ) * DELTA_HOURS
            errors[day_index, release_index, start:] = (
                data.net_load_kwh[day_index, start:] - net_hat[day_index, release_index, start:]
            )
            valid[day_index, release_index, start:] = valid_days[day_index]
    price_ratio = np.full_like(data.actual_price, np.nan)
    price_ratio[valid_days] = data.actual_price[valid_days] / price_hat[valid_days]
    return ForecastArchive(source, load_hat, price_hat, price_ratio, pv_hat, net_hat, errors, valid)


def build_scenarios(
    archive: ForecastArchive,
    day_index: int,
    release_index: int,
    start: int,
    mode: str,
) -> ScenarioSet:
    point_net = archive.net_kwh[day_index, release_index, start:]
    point_price = archive.price_hat[day_index, start:]
    if not np.isfinite(point_net).all() or not np.isfinite(point_price).all():
        raise ValueError(f"Point forecast unavailable on day index {day_index}, release {RELEASE_HOURS[release_index]}")
    if mode == "point":
        return ScenarioSet(point_net[None, :], point_price[None, :], np.ones(1), tuple(), True, "point_baseline")
    if mode != "paired":
        raise ValueError(f"Unknown scenario mode {mode!r}")

    candidates = range(max(1, day_index - SCENARIO_WINDOW_DAYS), day_index)
    eligible = [
        idx
        for idx in candidates
        if archive.valid[idx, release_index, start:].all()
        and np.isfinite(archive.price_ratio[idx, start:]).all()
    ]
    if len(eligible) < MIN_SCENARIO_DAYS:
        return ScenarioSet(
            point_net[None, :],
            point_price[None, :],
            np.ones(1),
            tuple(eligible),
            True,
            "insufficient_history",
        )
    index = np.asarray(eligible, dtype=int)
    scenario_net = point_net[None, :] + archive.errors_kwh[index, release_index, start:]
    scenario_price = point_price[None, :] * archive.price_ratio[index, start:]
    probabilities = np.full(len(index), 1.0 / len(index))
    if not np.isfinite(scenario_net).all() or not np.isfinite(scenario_price).all() or np.any(scenario_price <= 0):
        raise ValueError("Constructed scenarios contain invalid values")
    return ScenarioSet(scenario_net, scenario_price, probabilities, tuple(eligible), False, None)


def build_inherited_risk_margins(
    archive: ForecastArchive,
    question_variant: str,
    alpha: float | None,
) -> np.ndarray:
    """Reproduce the approved Q2/Q3 risk-margin rules without adding price uncertainty."""
    margins = np.zeros_like(archive.net_kwh)
    if alpha is None:
        return margins
    if question_variant == "Q4-2":
        for day_index in range(1, N_DAYS):
            prior = np.arange(max(1, day_index - 28), day_index, dtype=int)
            eligible = prior[np.all(archive.valid[prior, 0], axis=1)]
            if len(np.unique(eligible)) < 7:
                continue
            for group in range(6):
                section = slice(group * 24, (group + 1) * 24)
                samples = archive.errors_kwh[eligible, 0, section].reshape(-1)
                value = float(np.quantile(samples, alpha, method="linear"))
                margins[day_index, 0, section] = max(0.0, value)
        return margins
    if question_variant == "Q4-3":
        for day_index in range(1, N_DAYS):
            prior = np.arange(max(1, day_index - 28), day_index, dtype=int)
            if len(prior) < 7:
                continue
            for release_index, start in enumerate(RELEASE_STARTS):
                for interval in range(start, N_INTERVALS):
                    eligible = prior[archive.valid[prior, release_index, interval]]
                    if len(np.unique(eligible)) < 7:
                        continue
                    value = float(
                        np.quantile(
                            archive.errors_kwh[eligible, release_index, interval],
                            alpha,
                            method="linear",
                        )
                    )
                    margins[day_index, release_index, interval] = max(0.0, value)
        return margins
    raise ValueError(f"Unknown question variant {question_variant!r}")


def build_inherited_planning_case(
    data: Q4Data,
    archive: ForecastArchive,
    margins: np.ndarray,
    day_index: int,
    release_index: int,
    start: int,
) -> ScenarioSet:
    """为题给数据离线重算使用附件4目标日完整价格路径；不主张现实运行可预知未来实时价格。"""
    planned_net = archive.net_kwh[day_index, release_index, start:] + margins[
        day_index, release_index, start:
    ]
    price = data.actual_price[day_index, start:]
    if not np.isfinite(planned_net).all() or not np.isfinite(price).all() or np.any(price <= 0):
        raise ValueError(
            f"Inherited planning inputs unavailable on day {day_index}, release {RELEASE_HOURS[release_index]}"
        )
    return ScenarioSet(
        planned_net[None, :],
        price[None, :],
        np.ones(1),
        tuple(),
        False,
        "known_price_schedule_inherited_from_q2_q3",
    )


class ScenarioLP:
    def __init__(self) -> None:
        self.da_cache: dict[tuple[int, int], tuple[csr_matrix, list[tuple[float | None, float | None]], np.ndarray]] = {}
        self.tr_cache: dict[tuple[int, int], tuple[csr_matrix, list[tuple[float | None, float | None]], np.ndarray]] = {}

    def _da_structure(
        self, length: int, scenario_count: int
    ) -> tuple[csr_matrix, list[tuple[float | None, float | None]], np.ndarray]:
        key = (length, scenario_count)
        if key in self.da_cache:
            return self.da_cache[key]
        t_count = length
        j_count = scenario_count
        jt = j_count * t_count
        x0, c0, d0, s0, e0, w0 = 0, t_count, 2 * t_count, 3 * t_count, 4 * t_count, 4 * t_count + jt
        columns = 4 * t_count + 2 * jt
        matrix = lil_matrix((jt + t_count, columns), dtype=float)
        for j in range(j_count):
            for t in range(t_count):
                row = j * t_count + t
                matrix[row, [x0 + t, c0 + t, d0 + t, e0 + row, w0 + row]] = [1.0, -1.0, 1.0, 1.0, -1.0]
        for t in range(t_count):
            row = jt + t
            matrix[row, [c0 + t, d0 + t, s0 + t]] = [-ETA_CHARGE, 1.0 / ETA_DISCHARGE, 1.0]
            if t:
                matrix[row, s0 + t - 1] = -1.0
        bounds = (
            [(0.0, None)] * t_count
            + [(0.0, INTERVAL_LIMIT_KWH)] * t_count
            + [(0.0, INTERVAL_LIMIT_KWH)] * t_count
            + [(SOC_MIN_KWH, SOC_MAX_KWH)] * t_count
            + [(0.0, None)] * (2 * jt)
        )
        secondary = np.zeros(columns)
        secondary[c0 : c0 + t_count] = 1.0
        secondary[d0 : d0 + t_count] = 1.0
        value = (matrix.tocsr(), bounds, secondary)
        self.da_cache[key] = value
        return value

    def solve_da(self, scenarios: ScenarioSet, initial_soc: float, terminal_value: float) -> LPSolution:
        j_count, t_count = scenarios.net_kwh.shape
        jt = j_count * t_count
        matrix, bounds, secondary = self._da_structure(t_count, j_count)
        rhs = np.r_[scenarios.net_kwh.reshape(-1), np.zeros(t_count)]
        rhs[jt] = initial_soc
        x0, c0, d0, s0, e0, w0 = 0, t_count, 2 * t_count, 3 * t_count, 4 * t_count, 4 * t_count + jt
        mean_price = scenarios.probabilities @ scenarios.price
        primary = np.zeros(4 * t_count + 2 * jt)
        primary[x0 : x0 + t_count] = mean_price
        primary[s0 + t_count - 1] = -terminal_value
        primary[e0 : e0 + jt] = (5.0 * scenarios.probabilities[:, None] * scenarios.price).reshape(-1)
        return self._solve(
            "DA", matrix, rhs, bounds, primary, secondary, scenarios, initial_soc,
            (x0, c0, d0, s0, e0, w0), None,
        )

    def _tr_structure(
        self, length: int, scenario_count: int
    ) -> tuple[csr_matrix, list[tuple[float | None, float | None]], np.ndarray]:
        key = (length, scenario_count)
        if key in self.tr_cache:
            return self.tr_cache[key]
        t_count = length
        j_count = scenario_count
        jt = j_count * t_count
        y0, ap0, am0, c0, d0, s0, e0, w0 = (
            0, t_count, 2 * t_count, 3 * t_count, 4 * t_count, 5 * t_count, 6 * t_count, 6 * t_count + jt
        )
        columns = 6 * t_count + 2 * jt
        matrix = lil_matrix((jt + 2 * t_count, columns), dtype=float)
        for j in range(j_count):
            for t in range(t_count):
                row = j * t_count + t
                matrix[row, [y0 + t, c0 + t, d0 + t, e0 + row, w0 + row]] = [1.0, -1.0, 1.0, 1.0, -1.0]
        for t in range(t_count):
            row = jt + t
            matrix[row, [c0 + t, d0 + t, s0 + t]] = [-ETA_CHARGE, 1.0 / ETA_DISCHARGE, 1.0]
            if t:
                matrix[row, s0 + t - 1] = -1.0
            matrix[jt + t_count + t, [y0 + t, ap0 + t, am0 + t]] = [1.0, -1.0, 1.0]
        bounds = (
            [(0.0, None)] * (3 * t_count)
            + [(0.0, INTERVAL_LIMIT_KWH)] * t_count
            + [(0.0, INTERVAL_LIMIT_KWH)] * t_count
            + [(SOC_MIN_KWH, SOC_MAX_KWH)] * t_count
            + [(0.0, None)] * (2 * jt)
        )
        secondary = np.zeros(columns)
        secondary[c0 : c0 + t_count] = 1.0
        secondary[d0 : d0 + t_count] = 1.0
        value = (matrix.tocsr(), bounds, secondary)
        self.tr_cache[key] = value
        return value

    def solve_tr(
        self,
        scenarios: ScenarioSet,
        initial_soc: float,
        previous_contract: np.ndarray,
        terminal_value: float,
    ) -> LPSolution:
        j_count, t_count = scenarios.net_kwh.shape
        jt = j_count * t_count
        matrix, bounds, secondary = self._tr_structure(t_count, j_count)
        rhs = np.r_[scenarios.net_kwh.reshape(-1), np.zeros(t_count), previous_contract]
        rhs[jt] = initial_soc
        y0, ap0, am0, c0, d0, s0, e0, w0 = (
            0, t_count, 2 * t_count, 3 * t_count, 4 * t_count, 5 * t_count, 6 * t_count, 6 * t_count + jt
        )
        mean_price = scenarios.probabilities @ scenarios.price
        primary = np.zeros(6 * t_count + 2 * jt)
        primary[ap0 : ap0 + t_count] = 1.5 * mean_price
        primary[am0 : am0 + t_count] = -0.5 * mean_price
        primary[s0 + t_count - 1] = -terminal_value
        primary[e0 : e0 + jt] = (5.0 * scenarios.probabilities[:, None] * scenarios.price).reshape(-1)
        return self._solve(
            "TR", matrix, rhs, bounds, primary, secondary, scenarios, initial_soc,
            (y0, c0, d0, s0, e0, w0), previous_contract,
            adjustment_blocks=(ap0, am0),
        )

    @staticmethod
    def _solve(
        label: str,
        matrix: csr_matrix,
        rhs: np.ndarray,
        bounds: list[tuple[float | None, float | None]],
        primary: np.ndarray,
        secondary: np.ndarray,
        scenarios: ScenarioSet,
        initial_soc: float,
        block_starts: tuple[int, int, int, int, int, int],
        previous_contract: np.ndarray | None,
        adjustment_blocks: tuple[int, int] | None = None,
    ) -> LPSolution:
        stage1 = linprog(primary, A_eq=matrix, b_eq=rhs, bounds=bounds, method="highs")
        if not stage1.success:
            raise RuntimeError(f"{label} stage1 failed: {stage1.message}")
        stage2 = linprog(
            secondary,
            A_ub=csr_matrix(primary.reshape(1, -1)),
            b_ub=np.asarray([float(stage1.fun) + PRIMARY_TOLERANCE]),
            A_eq=matrix,
            b_eq=rhs,
            bounds=bounds,
            method="highs",
        )
        if not stage2.success:
            raise RuntimeError(f"{label} stage2 failed: {stage2.message}")
        t_count = scenarios.net_kwh.shape[1]
        j_count = scenarios.net_kwh.shape[0]
        x0, c0, d0, s0, e0, w0 = block_starts
        values = stage2.x
        contract = values[x0 : x0 + t_count]
        charge = values[c0 : c0 + t_count]
        discharge = values[d0 : d0 + t_count]
        soc = values[s0 : s0 + t_count]
        emergency = values[e0 : e0 + j_count * t_count].reshape(j_count, t_count)
        unused = values[w0 : w0 + j_count * t_count].reshape(j_count, t_count)
        balance = contract[None, :] + emergency + discharge[None, :] - scenarios.net_kwh - charge[None, :] - unused
        prior_soc = np.r_[initial_soc, soc[:-1]]
        storage = soc - prior_soc - ETA_CHARGE * charge + discharge / ETA_DISCHARGE
        gap = scenarios.net_kwh + charge[None, :] - discharge[None, :] - contract[None, :]
        emergency_formula = np.maximum(gap, 0.0)
        unused_formula = np.maximum(-gap, 0.0)
        emergency_residual = max(
            float(np.max(np.abs(emergency - emergency_formula))),
            float(np.max(np.abs(unused - unused_formula))),
            float(np.max(np.minimum(emergency, unused))),
        )
        adjustment_residual = 0.0
        if adjustment_blocks is not None and previous_contract is not None:
            ap0, am0 = adjustment_blocks
            increase = values[ap0 : ap0 + t_count]
            decrease = values[am0 : am0 + t_count]
            adjustment_residual = max(
                float(np.max(np.abs(contract - increase + decrease - previous_contract))),
                float(np.max(np.minimum(increase, decrease))),
            )
        return LPSolution(
            contract=contract,
            charge=charge,
            discharge=discharge,
            soc_end=soc,
            emergency=emergency,
            unused=unused,
            stage1_objective=float(stage1.fun),
            stage2_primary_objective=float(primary @ stage2.x),
            throughput=float(charge.sum() + discharge.sum()),
            max_balance_residual=float(np.max(np.abs(balance))),
            max_storage_residual=float(np.max(np.abs(storage))),
            max_emergency_formula_residual=emergency_residual,
            max_adjustment_residual=adjustment_residual,
            max_charge_discharge_overlap=float(np.max(np.minimum(charge, discharge))),
            stage1_status=str(stage1.message),
            stage2_status=str(stage2.message),
        )


def _fallback_da(point_net: np.ndarray, scenario_count: int, initial_soc: float) -> LPSolution:
    contract = np.maximum(point_net, 0.0)
    zeros = np.zeros_like(contract)
    gap = point_net - contract
    emergency = np.tile(np.maximum(gap, 0.0), (scenario_count, 1))
    unused = np.tile(np.maximum(-gap, 0.0), (scenario_count, 1))
    return LPSolution(
        contract, zeros.copy(), zeros.copy(), np.full_like(contract, initial_soc), emergency, unused,
        math.nan, math.nan, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0,
        "failed_fallback_activated", "failed_fallback_activated",
    )


def incremental_adjustment_fee(
    new_contract: np.ndarray, previous_contract: np.ndarray, actual_price: np.ndarray
) -> np.ndarray:
    increase = np.maximum(new_contract - previous_contract, 0.0)
    decrease = np.maximum(previous_contract - new_contract, 0.0)
    return actual_price * (1.5 * increase - 0.5 * decrease)


def run_policy(
    data: Q4Data,
    archive: ForecastArchive,
    config: PolicyConfig,
    margins: np.ndarray | None = None,
    max_days: int = N_DAYS,
    start_day_index: int = 0,
    initial_soc_override: float | None = None,
) -> StrategyRun:
    started = time_module.perf_counter()
    solver = ScenarioLP()
    day_limit = min(max_days, N_DAYS)
    if not 0 <= start_day_index < day_limit:
        raise ValueError("start_day_index must identify a day before max_days")
    release_indices = (0,) if config.question_variant == "Q4-2" else (0, 1, 2, 3)
    current_soc = SOC_INITIAL_KWH if initial_soc_override is None else float(initial_soc_override)
    if not SOC_MIN_KWH <= current_soc <= SOC_MAX_KWH:
        raise ValueError("Initial SOC must satisfy the storage energy bounds")
    run_initial_soc = current_soc
    daily_rows: list[dict[str, Any]] = []
    ledger_rows: list[dict[str, Any]] = []
    adjustment_rows: list[dict[str, Any]] = []
    plan_call_rows: list[dict[str, Any]] = []
    plan_trajectory_rows: list[dict[str, Any]] = []
    fallback_events: list[dict[str, Any]] = []
    max_checks: dict[str, float] = {}
    total_charge_full = 0.0
    total_discharge_full = 0.0
    plan_count = 0
    degenerate_count = 0

    if config.scenario_mode == "inherited" and margins is None:
        raise ValueError("Inherited Q2/Q3 execution requires an explicit risk-margin array")

    def planning_case(day_index: int, release_index: int, start: int) -> ScenarioSet:
        if config.scenario_mode == "inherited":
            assert margins is not None
            return build_inherited_planning_case(
                data, archive, margins, day_index, release_index, start
            )
        return build_scenarios(archive, day_index, release_index, start, config.scenario_mode)

    def update_max(name: str, value: float) -> None:
        if math.isfinite(value):
            max_checks[name] = max(max_checks.get(name, 0.0), float(value))

    def record_solution(
        day_index: int,
        release_index: int,
        start: int,
        scenarios: ScenarioSet,
        solution: LPSolution,
        call_id: str,
        terminal_value: float,
    ) -> None:
        nonlocal plan_count, degenerate_count
        plan_count += 1
        degenerate_count += int(scenarios.degenerate)
        update_max("planning_balance_max_abs", solution.max_balance_residual)
        update_max("planning_storage_recurrence_max_abs", solution.max_storage_residual)
        update_max("scenario_emergency_formula_max_abs", solution.max_emergency_formula_residual)
        update_max("planning_adjustment_max_abs", solution.max_adjustment_residual)
        update_max("planning_charge_discharge_overlap_max", solution.max_charge_discharge_overlap)
        update_max("stage2_objective_excess_yuan", max(0.0, solution.stage2_primary_objective - solution.stage1_objective))
        if not config.collect_evidence:
            return
        historical_dates = [data.dates[index].isoformat() for index in scenarios.historical_day_indices]
        plan_call_rows.append(
            {
                "strategy_id": config.strategy_id,
                "call_id": call_id,
                "date": data.dates[day_index].isoformat(),
                "release_hour": RELEASE_HOURS[release_index],
                "start_interval_index": start + 1,
                "remaining_interval_count": len(solution.contract),
                "scenario_count": len(scenarios.probabilities),
                "scenario_mode": config.scenario_mode,
                "degenerate": scenarios.degenerate,
                "degenerate_reason": scenarios.reason,
                "historical_scenario_dates": ";".join(historical_dates),
                "stage1_objective_yuan": solution.stage1_objective,
                "stage2_primary_objective_yuan": solution.stage2_primary_objective,
                "throughput_kwh": solution.throughput,
                "terminal_value_yuan_per_kwh": terminal_value,
                "stage1_status": solution.stage1_status,
                "stage2_status": solution.stage2_status,
            }
        )
        expected_emergency = scenarios.probabilities @ solution.emergency
        expected_unused = scenarios.probabilities @ solution.unused
        for local, interval in enumerate(range(start, N_INTERVALS)):
            plan_trajectory_rows.append(
                {
                    "strategy_id": config.strategy_id,
                    "call_id": call_id,
                    "date": data.dates[day_index].isoformat(),
                    "release_hour": RELEASE_HOURS[release_index],
                    "interval_index": interval + 1,
                    "interval_label": data.interval_labels[interval],
                    "contract_kwh": solution.contract[local],
                    "reference_charge_kwh": solution.charge[local],
                    "reference_discharge_kwh": solution.discharge[local],
                    "reference_soc_end_kwh": solution.soc_end[local],
                    "expected_emergency_kwh": expected_emergency[local],
                    "expected_unused_kwh": expected_unused[local],
                }
            )

    for day_index in range(start_day_index, day_limit):
        day = data.dates[day_index]
        day_initial_soc = current_soc
        q = np.zeros(N_INTERVALS)
        final_contract = np.zeros(N_INTERVALS)
        current_plan_charge = np.zeros(N_INTERVALS)
        current_plan_discharge = np.zeros(N_INTERVALS)
        current_plan_emergency = np.zeros(N_INTERVALS)
        current_plan_unused = np.zeros(N_INTERVALS)
        actual_charge = np.zeros(N_INTERVALS)
        actual_discharge = np.zeros(N_INTERVALS)
        actual_emergency = np.zeros(N_INTERVALS)
        actual_unused = np.zeros(N_INTERVALS)
        soc_end = np.zeros(N_INTERVALS)
        adjustment_fee_by_interval = np.zeros(N_INTERVALS)
        adjustment_increase_by_interval = np.zeros(N_INTERVALS)
        adjustment_decrease_by_interval = np.zeros(N_INTERVALS)
        adjustment_count_by_interval = np.zeros(N_INTERVALS, dtype=int)
        executing_release = np.zeros(N_INTERVALS, dtype=int)
        day_scenario_counts: list[int] = []

        if day_index == 0:
            terminal_value = 0.0
        else:
            da_scenarios = planning_case(day_index, 0, 0)
            terminal_value = (
                0.0
                if day_index == FINAL_DAY_INDEX
                else GAMMA * ETA_DISCHARGE * float(np.median(da_scenarios.probabilities @ da_scenarios.price))
            )
            try:
                da = solver.solve_da(da_scenarios, current_soc, terminal_value)
            except RuntimeError as exc:
                da = _fallback_da(archive.net_kwh[day_index, 0], len(da_scenarios.probabilities), current_soc)
                fallback_events.append({"date": day.isoformat(), "release_hour": 0, "error": str(exc)})
            q[:] = da.contract
            final_contract[:] = q
            current_plan_charge[:] = da.charge
            current_plan_discharge[:] = da.discharge
            current_plan_emergency[:] = da_scenarios.probabilities @ da.emergency
            current_plan_unused[:] = da_scenarios.probabilities @ da.unused
            day_scenario_counts.append(len(da_scenarios.probabilities))
            record_solution(day_index, 0, 0, da_scenarios, da, f"{day.isoformat()}_00", terminal_value)

        for position, release_index in enumerate(release_indices):
            release_hour = RELEASE_HOURS[release_index]
            start = RELEASE_STARTS[release_index]
            if day_index > 0 and release_index > 0:
                scenarios = planning_case(day_index, release_index, start)
                previous = final_contract[start:].copy()
                try:
                    solution = solver.solve_tr(scenarios, current_soc, previous, terminal_value)
                    new_contract = solution.contract
                except RuntimeError as exc:
                    fallback_events.append({"date": day.isoformat(), "release_hour": release_hour, "error": str(exc)})
                    new_contract = previous.copy()
                    point_gap = archive.net_kwh[day_index, release_index, start:] - new_contract
                    solution = LPSolution(
                        new_contract,
                        np.zeros_like(new_contract),
                        np.zeros_like(new_contract),
                        np.full_like(new_contract, current_soc),
                        np.tile(np.maximum(point_gap, 0.0), (len(scenarios.probabilities), 1)),
                        np.tile(np.maximum(-point_gap, 0.0), (len(scenarios.probabilities), 1)),
                        math.nan,
                        math.nan,
                        0.0,
                        0.0,
                        0.0,
                        0.0,
                        0.0,
                        0.0,
                        "failed_hold_previous",
                        "failed_hold_previous",
                    )
                increase = np.maximum(new_contract - previous, 0.0)
                decrease = np.maximum(previous - new_contract, 0.0)
                incremental_fee = incremental_adjustment_fee(new_contract, previous, data.actual_price[day_index, start:])
                changed = np.abs(new_contract - previous) > EVENT_TOLERANCE
                adjustment_fee_by_interval[start:] += incremental_fee
                adjustment_increase_by_interval[start:] += increase
                adjustment_decrease_by_interval[start:] += decrease
                adjustment_count_by_interval[start:] += changed.astype(int)
                final_contract[start:] = new_contract
                current_plan_charge[start:] = solution.charge
                current_plan_discharge[start:] = solution.discharge
                current_plan_emergency[start:] = scenarios.probabilities @ solution.emergency
                current_plan_unused[start:] = scenarios.probabilities @ solution.unused
                executing_release[start:] = release_hour
                day_scenario_counts.append(len(scenarios.probabilities))
                call_id = f"{day.isoformat()}_{release_hour:02d}"
                record_solution(day_index, release_index, start, scenarios, solution, call_id, terminal_value)
                if config.collect_evidence:
                    for local, interval in enumerate(range(start, N_INTERVALS)):
                        adjustment_rows.append(
                            {
                                "strategy_id": config.strategy_id,
                                "date": day.isoformat(),
                                "release_hour": release_hour,
                                "interval_index": interval + 1,
                                "interval_label": data.interval_labels[interval],
                                "actual_price_yuan_per_kwh": data.actual_price[day_index, interval],
                                "previous_contract_kwh": previous[local],
                                "new_contract_kwh": new_contract[local],
                                "increase_kwh": increase[local],
                                "decrease_kwh": decrease[local],
                                "incremental_adjustment_fee_yuan": incremental_fee[local],
                                "changed": bool(changed[local]),
                            }
                        )

            end = RELEASE_STARTS[release_indices[position + 1]] if position + 1 < len(release_indices) else N_INTERVALS
            for interval in range(start, end):
                net_gap = data.net_load_kwh[day_index, interval] - final_contract[interval]
                if day_index == 0:
                    charge = 0.0
                    discharge = 0.0
                else:
                    charge = min(
                        max(-net_gap, 0.0),
                        INTERVAL_LIMIT_KWH,
                        max(0.0, (SOC_MAX_KWH - current_soc) / ETA_CHARGE),
                    )
                    discharge = min(
                        max(net_gap, 0.0),
                        INTERVAL_LIMIT_KWH,
                        max(0.0, ETA_DISCHARGE * (current_soc - SOC_MIN_KWH)),
                    )
                residual = net_gap + charge - discharge
                emergency = max(residual, 0.0)
                unused = max(-residual, 0.0)
                next_soc = current_soc + ETA_CHARGE * charge - discharge / ETA_DISCHARGE
                actual_charge[interval] = charge
                actual_discharge[interval] = discharge
                actual_emergency[interval] = emergency
                actual_unused[interval] = unused
                soc_end[interval] = next_soc
                total_charge_full += charge
                total_discharge_full += discharge
                update_max("actual_soc_lower_violation", max(0.0, SOC_MIN_KWH - next_soc))
                update_max("actual_soc_upper_violation", max(0.0, next_soc - SOC_MAX_KWH))
                update_max("actual_charge_limit_violation", max(0.0, charge - INTERVAL_LIMIT_KWH))
                update_max("actual_discharge_limit_violation", max(0.0, discharge - INTERVAL_LIMIT_KWH))
                update_max("actual_charge_discharge_overlap", min(charge, discharge))
                update_max("actual_emergency_unused_overlap", min(emergency, unused))
                balance = (
                    final_contract[interval]
                    + emergency
                    + data.pv_kw[day_index, interval] * DELTA_HOURS
                    + discharge
                    - data.load_kw[day_index, interval] * DELTA_HOURS
                    - charge
                    - unused
                )
                recurrence = next_soc - current_soc - ETA_CHARGE * charge + discharge / ETA_DISCHARGE
                update_max("actual_balance_max_abs", abs(balance))
                update_max("actual_storage_recurrence_max_abs", abs(recurrence))
                current_soc = next_soc

        original_cost_by_interval = data.actual_price[day_index] * q
        contract_fee_by_interval = original_cost_by_interval + adjustment_fee_by_interval
        path_identity = contract_fee_by_interval - data.actual_price[day_index] * (
            final_contract + 0.5 * (adjustment_increase_by_interval + adjustment_decrease_by_interval)
        )
        update_max("contract_path_identity_max_abs", float(np.max(np.abs(path_identity))))
        emergency_cost_by_interval = 5.0 * data.actual_price[day_index] * actual_emergency
        total_cost_by_interval = contract_fee_by_interval + emergency_cost_by_interval
        daily_rows.append(
            {
                "strategy_id": config.strategy_id,
                "question_variant": config.question_variant,
                "date": day.isoformat(),
                "initial_soc_kwh": day_initial_soc,
                "terminal_soc_kwh": current_soc,
                "original_contract_kwh": float(q.sum()),
                "final_contract_kwh": float(final_contract.sum()),
                "adjustment_increase_kwh": float(adjustment_increase_by_interval.sum()),
                "adjustment_decrease_kwh": float(adjustment_decrease_by_interval.sum()),
                "adjustment_transaction_count": int(adjustment_count_by_interval.sum()),
                "emergency_purchase_kwh": float(actual_emergency.sum()),
                "actual_total_purchase_kwh": float(final_contract.sum() + actual_emergency.sum()),
                "charge_kwh": float(actual_charge.sum()),
                "discharge_kwh": float(actual_discharge.sum()),
                "unused_kwh": float(actual_unused.sum()),
                "original_contract_cost_yuan": float(original_cost_by_interval.sum()),
                "adjustment_net_cost_yuan": float(adjustment_fee_by_interval.sum()),
                "contract_total_cost_yuan": float(contract_fee_by_interval.sum()),
                "emergency_cost_yuan": float(emergency_cost_by_interval.sum()),
                "actual_total_cost_yuan": float(total_cost_by_interval.sum()),
                "planning_call_count": len(day_scenario_counts),
                "min_scenario_count": min(day_scenario_counts) if day_scenario_counts else 0,
                "max_scenario_count": max(day_scenario_counts) if day_scenario_counts else 0,
            }
        )
        if config.collect_evidence:
            prior_soc = np.r_[day_initial_soc, soc_end[:-1]]
            for interval in range(N_INTERVALS):
                release_index = RELEASE_HOURS.index(int(executing_release[interval]))
                ledger_rows.append(
                    {
                        "strategy_id": config.strategy_id,
                        "question_variant": config.question_variant,
                        "date": day.isoformat(),
                        "interval_index": interval + 1,
                        "interval_start": data.interval_starts[interval],
                        "interval_end": data.interval_ends[interval],
                        "interval_label": data.interval_labels[interval],
                        "executing_release_hour": int(executing_release[interval]),
                        "actual_price_yuan_per_kwh": data.actual_price[day_index, interval],
                        "actual_load_kw": data.load_kw[day_index, interval],
                        "actual_pv_kw": data.pv_kw[day_index, interval],
                        "actual_net_load_kwh": data.net_load_kwh[day_index, interval],
                        "forecast_load_kw": archive.load_kw[day_index, interval] if day_index else np.nan,
                        "forecast_pv_kw": archive.pv_kw[day_index, release_index, interval] if day_index else np.nan,
                        "forecast_net_load_kwh": archive.net_kwh[day_index, release_index, interval] if day_index else np.nan,
                        "forecast_price_yuan_per_kwh": (
                            data.actual_price[day_index, interval]
                            if day_index and config.scenario_mode == "inherited"
                            else archive.price_hat[day_index, interval] if day_index else np.nan
                        ),
                        "original_contract_kwh": q[interval],
                        "final_contract_kwh": final_contract[interval],
                        "adjustment_increase_kwh": adjustment_increase_by_interval[interval],
                        "adjustment_decrease_kwh": adjustment_decrease_by_interval[interval],
                        "adjustment_transaction_count": int(adjustment_count_by_interval[interval]),
                        "reference_charge_kwh": current_plan_charge[interval],
                        "reference_discharge_kwh": current_plan_discharge[interval],
                        "reference_expected_emergency_kwh": current_plan_emergency[interval],
                        "reference_expected_unused_kwh": current_plan_unused[interval],
                        "actual_charge_kwh": actual_charge[interval],
                        "actual_discharge_kwh": actual_discharge[interval],
                        "actual_emergency_kwh": actual_emergency[interval],
                        "actual_unused_kwh": actual_unused[interval],
                        "soc_start_kwh": prior_soc[interval],
                        "soc_end_kwh": soc_end[interval],
                        "original_contract_cost_yuan": original_cost_by_interval[interval],
                        "cumulative_adjustment_fee_yuan": adjustment_fee_by_interval[interval],
                        "contract_total_cost_yuan": contract_fee_by_interval[interval],
                        "emergency_cost_yuan": emergency_cost_by_interval[interval],
                        "actual_total_cost_yuan": total_cost_by_interval[interval],
                    }
                )

    daily = pd.DataFrame(daily_rows)
    evaluation = daily[daily["date"] >= "2025-02-01"].copy()
    if evaluation.empty:
        evaluation = daily.copy()
    metrics = {
        "strategy_id": config.strategy_id,
        "role": config.role,
        "question_variant": config.question_variant,
        "forecast_source": config.forecast_source,
        "scenario_mode": config.scenario_mode,
        "alpha": config.alpha,
        "price_information": (
            FORMAL_PRICE_INFORMATION
            if config.scenario_mode == "inherited"
            else "extension_price_forecast_or_joint_historical_scenarios_not_formal_q4_round2"
        ),
        "simulation_start": daily.iloc[0]["date"],
        "simulation_days": len(daily),
        "run_initial_soc_kwh": run_initial_soc,
        "evaluation_start": evaluation.iloc[0]["date"],
        "evaluation_end": evaluation.iloc[-1]["date"],
        "evaluation_days": len(evaluation),
        "scenario_window_days": SCENARIO_WINDOW_DAYS,
        "minimum_scenario_days": MIN_SCENARIO_DAYS,
        "gamma": GAMMA,
        "initial_soc_evaluation_kwh": float(evaluation.iloc[0]["initial_soc_kwh"]),
        "terminal_soc_evaluation_kwh": float(evaluation.iloc[-1]["terminal_soc_kwh"]),
        "original_contract_kwh": float(evaluation["original_contract_kwh"].sum()),
        "final_contract_kwh": float(evaluation["final_contract_kwh"].sum()),
        "adjustment_increase_kwh": float(evaluation["adjustment_increase_kwh"].sum()),
        "adjustment_decrease_kwh": float(evaluation["adjustment_decrease_kwh"].sum()),
        "adjustment_transaction_count": int(evaluation["adjustment_transaction_count"].sum()),
        "emergency_purchase_kwh": float(evaluation["emergency_purchase_kwh"].sum()),
        "actual_total_purchase_kwh": float(evaluation["actual_total_purchase_kwh"].sum()),
        "charge_kwh": float(evaluation["charge_kwh"].sum()),
        "discharge_kwh": float(evaluation["discharge_kwh"].sum()),
        "original_contract_cost_yuan": float(evaluation["original_contract_cost_yuan"].sum()),
        "adjustment_net_cost_yuan": float(evaluation["adjustment_net_cost_yuan"].sum()),
        "contract_total_cost_yuan": float(evaluation["contract_total_cost_yuan"].sum()),
        "emergency_cost_yuan": float(evaluation["emergency_cost_yuan"].sum()),
        "actual_total_cost_yuan": float(evaluation["actual_total_cost_yuan"].sum()),
        "planning_calls_full_run": plan_count,
        "degenerate_scenario_calls_full_run": degenerate_count,
        "fallback_event_count": len(fallback_events),
    }
    fee_reconciliation = abs(
        metrics["contract_total_cost_yuan"]
        - metrics["original_contract_cost_yuan"]
        - metrics["adjustment_net_cost_yuan"]
    )
    total_reconciliation = abs(
        metrics["actual_total_cost_yuan"]
        - metrics["contract_total_cost_yuan"]
        - metrics["emergency_cost_yuan"]
    )
    cross_day = np.abs(daily["initial_soc_kwh"].to_numpy()[1:] - daily["terminal_soc_kwh"].to_numpy()[:-1])
    energy_identity = total_discharge_full - (
        ETA_CHARGE * ETA_DISCHARGE * total_charge_full
        - ETA_DISCHARGE * (current_soc - run_initial_soc)
    )
    validation = {
        **max_checks,
        "fee_reconciliation_abs_yuan": float(fee_reconciliation),
        "total_cost_reconciliation_abs_yuan": float(total_reconciliation),
        "cross_day_soc_max_abs_kwh": float(np.max(cross_day)) if len(cross_day) else 0.0,
        "full_run_storage_energy_identity_abs_kwh": abs(float(energy_identity)),
        "incremental_example_100_to_80_to_100_yuan": float(
            100.0
            + incremental_adjustment_fee(np.array([80.0]), np.array([100.0]), np.array([1.0]))[0]
            + incremental_adjustment_fee(np.array([100.0]), np.array([80.0]), np.array([1.0]))[0]
        ),
        "fallback_events": fallback_events,
    }
    planning_values = [
        validation.get("planning_balance_max_abs", 0.0),
        validation.get("planning_storage_recurrence_max_abs", 0.0),
        validation.get("scenario_emergency_formula_max_abs", 0.0),
        validation.get("planning_adjustment_max_abs", 0.0),
        validation.get("planning_charge_discharge_overlap_max", 0.0),
    ]
    physical_values = [
        validation.get("actual_soc_lower_violation", 0.0),
        validation.get("actual_soc_upper_violation", 0.0),
        validation.get("actual_charge_limit_violation", 0.0),
        validation.get("actual_discharge_limit_violation", 0.0),
        validation.get("actual_charge_discharge_overlap", 0.0),
        validation.get("actual_emergency_unused_overlap", 0.0),
        validation.get("actual_balance_max_abs", 0.0),
        validation.get("actual_storage_recurrence_max_abs", 0.0),
        validation.get("contract_path_identity_max_abs", 0.0),
        validation["fee_reconciliation_abs_yuan"],
        validation["total_cost_reconciliation_abs_yuan"],
        validation["cross_day_soc_max_abs_kwh"],
        validation["full_run_storage_energy_identity_abs_kwh"],
    ]
    validation["planning_checks_pass"] = bool(max(planning_values, default=0.0) <= PLANNING_CHECK_TOLERANCE)
    validation["physical_and_cost_checks_pass"] = bool(max(physical_values, default=0.0) <= PHYSICAL_CHECK_TOLERANCE)
    validation["hard_checks_pass"] = bool(
        validation["planning_checks_pass"]
        and validation["physical_and_cost_checks_pass"]
        and validation["incremental_example_100_to_80_to_100_yuan"] == 120.0
        and not fallback_events
    )
    return StrategyRun(
        config=config,
        runtime_seconds=time_module.perf_counter() - started,
        metrics=metrics,
        validation=validation,
        daily=daily,
        ledger=pd.DataFrame(ledger_rows) if config.collect_evidence else None,
        adjustments=(
            pd.DataFrame(adjustment_rows)
            if adjustment_rows
            else pd.DataFrame(
                columns=[
                    "strategy_id", "date", "release_hour", "interval_index", "interval_label",
                    "actual_price_yuan_per_kwh", "previous_contract_kwh", "new_contract_kwh",
                    "increase_kwh", "decrease_kwh", "incremental_adjustment_fee_yuan", "changed",
                ]
            )
        ) if config.collect_evidence else None,
        plan_calls=pd.DataFrame(plan_call_rows) if config.collect_evidence else None,
        plan_trajectory=pd.DataFrame(plan_trajectory_rows) if config.collect_evidence else None,
    )


def emergency_events(ledger: pd.DataFrame) -> pd.DataFrame:
    events: list[dict[str, Any]] = []
    for day, group in ledger.groupby("date", sort=True):
        ordered = group.sort_values("interval_index").reset_index(drop=True)
        amounts = ordered["actual_emergency_kwh"].to_numpy(dtype=float)
        active = amounts > EVENT_TOLERANCE
        cursor = 0
        event_index = 0
        while cursor < N_INTERVALS:
            if not active[cursor]:
                cursor += 1
                continue
            start = cursor
            while cursor + 1 < N_INTERVALS and active[cursor + 1]:
                cursor += 1
            end = cursor
            event_index += 1
            events.append(
                {
                    "date": day,
                    "event_index": event_index,
                    "start_interval_index": start + 1,
                    "end_interval_index": end + 1,
                    "emergency_time_range": f"{ordered.loc[start, 'interval_start']}-{ordered.loc[end, 'interval_end']}",
                    "emergency_purchase_kwh": float(amounts[start : end + 1].sum()),
                }
            )
            cursor += 1
    return pd.DataFrame(events)


def storage_summary(ledger: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for day, group in ledger.groupby("date", sort=True):
        ordered = group.sort_values("interval_index").reset_index(drop=True)
        for block in range(6):
            start = block * 24
            end = start + 24
            rows.append(
                {
                    "date": day,
                    "block_index": block + 1,
                    "time_range": f"{_fmt_minutes(block * 240)}-{_fmt_minutes((block + 1) * 240)}",
                    "charge_kwh": float(ordered.loc[start : end - 1, "actual_charge_kwh"].sum()),
                    "discharge_kwh": float(ordered.loc[start : end - 1, "actual_discharge_kwh"].sum()),
                    "soc_00_kwh": float(ordered.loc[0, "soc_start_kwh"]) if block == 0 else np.nan,
                    "soc_24_kwh": float(ordered.loc[N_INTERVALS - 1, "soc_end_kwh"]) if block == 5 else np.nan,
                }
            )
    return pd.DataFrame(rows)


def specified_tables(
    ledger: pd.DataFrame, storage: pd.DataFrame, events: pd.DataFrame
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    selected_dates = {"2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"}
    table1 = ledger[ledger["date"].isin(selected_dates)][
        [
            "date", "interval_index", "interval_label", "original_contract_kwh", "final_contract_kwh",
            "actual_emergency_kwh", "actual_total_cost_yuan",
        ]
    ].copy()
    table2 = storage[storage["date"].isin(selected_dates)].copy()
    table3 = events[events["date"].isin(selected_dates)].copy()
    return table1, table2, table3


def forecast_diagnostics(data: Q4Data, archive: ForecastArchive) -> dict[str, Any]:
    price_mask = np.isfinite(archive.price_hat)
    diagnostics: dict[str, Any] = {
        "archive_id": archive.archive_id,
        "price_mae_yuan_per_kwh": float(np.mean(np.abs(data.actual_price[price_mask] - archive.price_hat[price_mask]))),
        "release_metrics": {},
    }
    for release_index, (hour, start) in enumerate(zip(RELEASE_HOURS, RELEASE_STARTS)):
        mask = archive.valid[:, release_index, start:]
        errors = archive.errors_kwh[:, release_index, start:][mask]
        price_residual = (archive.price_ratio[:, start:] - 1.0)[mask]
        correlation = float(np.corrcoef(errors, price_residual)[0, 1]) if len(errors) > 1 else None
        diagnostics["release_metrics"][str(hour)] = {
            "observations": int(len(errors)),
            "net_load_mae_kwh": float(np.mean(np.abs(errors))),
            "net_error_price_ratio_residual_correlation": correlation,
        }
    return diagnostics


def environment_info() -> dict[str, Any]:
    return {
        "python": sys.version,
        "platform": platform.platform(),
        "numpy": np.__version__,
        "pandas": pd.__version__,
        "scipy": scipy.__version__,
        "solver": "scipy.optimize.linprog(method='highs')",
        "random_seed": RANDOM_SEED,
    }


def _configurations(suite: str) -> list[PolicyConfig]:
    if suite == "extension":
        return [
            PolicyConfig("M4J_Q42_paired", "extension", "Q4-2", "historical", "paired", False, alpha=None, collect_evidence=True),
            PolicyConfig("B4P_Q42_point", "extension_baseline", "Q4-2", "historical", "point", False, alpha=None),
            PolicyConfig("M4J_Q43_paired", "extension", "Q4-3", "official", "paired", True, alpha=None, collect_evidence=True),
            PolicyConfig("B4P_Q43_point", "extension_baseline", "Q4-3", "official", "point", True, alpha=None),
        ]
    return [
        PolicyConfig("M4I_Q42_inherit_M3_alpha08", "main", "Q4-2", "historical", "inherited", False, alpha=0.8, collect_evidence=True),
        PolicyConfig("B4I_Q42_inherit_M3_b0", "usable_baseline", "Q4-2", "historical", "inherited", False, alpha=None),
        PolicyConfig("M4I_Q43_inherit_M1I_alpha08", "main", "Q4-3", "official", "inherited", True, alpha=0.8, collect_evidence=True),
        PolicyConfig("B4I_Q43_inherit_M1I_b0", "usable_baseline", "Q4-3", "official", "inherited", True, alpha=None),
    ]


def run_suite(suite: str, max_days: int | None = None) -> None:
    data = load_data()
    archives = {
        "historical": build_forecast_archive(data, "historical"),
        "official": build_forecast_archive(data, "official"),
    }
    effective_days = max_days if max_days is not None else (60 if suite == "pilot" else N_DAYS)
    output_round = (
        "pilot_inherited_price"
        if suite == "pilot"
        else "extension_rerun" if suite == "extension" else "round2_inherited_price"
    )
    round_dir = PROJECT_ROOT / "results" / "Q4" / "experiments" / output_round
    table_dir = round_dir / "tables"
    metric_dir = round_dir / "metrics"
    table_dir.mkdir(parents=True, exist_ok=True)
    metric_dir.mkdir(parents=True, exist_ok=True)

    runs: list[StrategyRun] = []
    margin_cache: dict[tuple[str, float | None], np.ndarray] = {}
    for config in _configurations(suite):
        print(f"START {config.strategy_id} days={effective_days}", flush=True)
        margins = None
        if config.scenario_mode == "inherited":
            key = (config.question_variant, config.alpha)
            if key not in margin_cache:
                margin_cache[key] = build_inherited_risk_margins(
                    archives[config.forecast_source], config.question_variant, config.alpha
                )
            margins = margin_cache[key]
        run = run_policy(
            data,
            archives[config.forecast_source],
            config,
            margins=margins,
            max_days=effective_days,
        )
        print(
            f"DONE {config.strategy_id} cost={run.metrics['actual_total_cost_yuan']:.2f} "
            f"emergency={run.metrics['emergency_purchase_kwh']:.3f} "
            f"plans={run.metrics['planning_calls_full_run']} runtime={run.runtime_seconds:.2f}s",
            flush=True,
        )
        if not run.validation["hard_checks_pass"]:
            raise RuntimeError(f"Hard checks failed for {config.strategy_id}: {run.validation}")
        runs.append(run)

    by_id = {run.config.strategy_id: run for run in runs}
    comparison_rows: list[dict[str, Any]] = []
    comparison_ids = (
        [("Q4-2", "M4J_Q42_paired", "B4P_Q42_point"), ("Q4-3", "M4J_Q43_paired", "B4P_Q43_point")]
        if suite == "extension"
        else [
            ("Q4-2", "M4I_Q42_inherit_M3_alpha08", "B4I_Q42_inherit_M3_b0"),
            ("Q4-3", "M4I_Q43_inherit_M1I_alpha08", "B4I_Q43_inherit_M1I_b0"),
        ]
    )
    for variant, main_id, baseline_id in comparison_ids:
        main_run = by_id[main_id]
        baseline_run = by_id[baseline_id]
        for run in (main_run, baseline_run):
            comparison_rows.append(
                {
                    "question_variant": variant,
                    "strategy_id": run.config.strategy_id,
                    "role": run.config.role,
                    "scenario_mode": run.config.scenario_mode,
                    "actual_total_cost_yuan": run.metrics["actual_total_cost_yuan"],
                    "contract_total_cost_yuan": run.metrics["contract_total_cost_yuan"],
                    "emergency_cost_yuan": run.metrics["emergency_cost_yuan"],
                    "emergency_purchase_kwh": run.metrics["emergency_purchase_kwh"],
                    "final_contract_kwh": run.metrics["final_contract_kwh"],
                    "initial_soc_evaluation_kwh": run.metrics["initial_soc_evaluation_kwh"],
                    "terminal_soc_evaluation_kwh": run.metrics["terminal_soc_evaluation_kwh"],
                    "runtime_seconds": run.runtime_seconds,
                    "cost_change_vs_baseline_pct": (
                        run.metrics["actual_total_cost_yuan"] / baseline_run.metrics["actual_total_cost_yuan"] - 1.0
                    ) * 100.0,
                }
            )
    save_frame(pd.DataFrame(comparison_rows), table_dir / "strategy_comparison.csv")

    output_files = ["tables/strategy_comparison.csv"]
    for run in runs:
        prefix = run.config.question_variant.lower().replace("-", "") + ("_main" if run.config.role == "main" else "_baseline")
        evaluation_daily = run.daily[run.daily["date"] >= "2025-02-01"].copy()
        save_frame(evaluation_daily, table_dir / f"{prefix}_daily_summary.csv")
        output_files.append(f"tables/{prefix}_daily_summary.csv")
        write_json(metric_dir / f"{prefix}_metrics.json", run.metrics)
        write_json(metric_dir / f"{prefix}_validation.json", run.validation)
        output_files.extend([f"metrics/{prefix}_metrics.json", f"metrics/{prefix}_validation.json"])
        if run.config.collect_evidence:
            assert run.ledger is not None and run.adjustments is not None
            assert run.plan_calls is not None and run.plan_trajectory is not None
            evaluation_ledger = run.ledger[run.ledger["date"] >= "2025-02-01"].copy()
            evaluation_adjustments = run.adjustments[run.adjustments["date"] >= "2025-02-01"].copy()
            events = emergency_events(evaluation_ledger)
            storage = storage_summary(evaluation_ledger)
            table1, table2, table3 = specified_tables(evaluation_ledger, storage, events)
            frames = {
                f"{prefix}_interval_ledger.csv": evaluation_ledger,
                f"{prefix}_adjustment_events.csv": evaluation_adjustments,
                f"{prefix}_planning_calls.csv": run.plan_calls,
                f"{prefix}_planning_trajectory.csv": run.plan_trajectory,
                f"{prefix}_emergency_events.csv": events,
                f"{prefix}_storage_4h.csv": storage,
                f"{prefix}_specified_date_table1.csv": table1,
                f"{prefix}_specified_date_table2.csv": table2,
                f"{prefix}_specified_date_table3.csv": table3,
            }
            for name, frame in frames.items():
                save_frame(frame, table_dir / name)
                output_files.append(f"tables/{name}")

    write_json(metric_dir / "historical_forecast_diagnostics.json", forecast_diagnostics(data, archives["historical"]))
    write_json(metric_dir / "official_forecast_diagnostics.json", forecast_diagnostics(data, archives["official"]))
    output_files.extend(["metrics/historical_forecast_diagnostics.json", "metrics/official_forecast_diagnostics.json"])

    q42_main = by_id[comparison_ids[0][1]]
    q42_base = by_id[comparison_ids[0][2]]
    q43_main = by_id[comparison_ids[1][1]]
    q43_base = by_id[comparison_ids[1][2]]
    degeneracy: dict[str, Any] = {}
    for key, run in [("q42", q42_main), ("q43", q43_main)]:
        assert run.ledger is not None
        eval_ledger = run.ledger[run.ledger["date"] >= "2025-02-01"]
        q_profiles = eval_ledger.pivot(index="date", columns="interval_index", values="original_contract_kwh")
        y_profiles = eval_ledger.pivot(index="date", columns="interval_index", values="final_contract_kwh")
        events = emergency_events(eval_ledger)
        degeneracy[key] = {
            "unique_original_contract_profiles_rounded_6dp": int(len(np.unique(np.round(q_profiles.to_numpy(), 6), axis=0))),
            "unique_final_contract_profiles_rounded_6dp": int(len(np.unique(np.round(y_profiles.to_numpy(), 6), axis=0))),
            "nonzero_final_vs_original_intervals": int(np.count_nonzero(np.abs(eval_ledger["final_contract_kwh"] - eval_ledger["original_contract_kwh"]) > EVENT_TOLERANCE)),
            "emergency_event_count": int(len(events)),
        }
    run_summary = {
        "schema_version": 1,
        "question": "Q4",
        "round": output_round,
        "implementation_target": "python",
        "random_seed": RANDOM_SEED,
        "approved_decision_id": (
            "q4_method_choice_joint_scenario_v1"
            if suite == "extension"
            else "q4_method_choice_inherit_q2_q3_variable_price_v2"
        ),
        "suite": suite,
        "days_executed": effective_days,
        "methods": [
            {
                "method_id": run.config.strategy_id,
                "role": run.config.role,
                "script": "code/Q4/q4_model.py",
                "status": "success",
                "execution_time_seconds": run.runtime_seconds,
                "input_files": [
                    str(ACTUAL_PATH.relative_to(PROJECT_ROOT.parent)),
                    *(
                        [str(FORECAST_PATH.relative_to(PROJECT_ROOT.parent))]
                        if run.config.question_variant == "Q4-3"
                        else []
                    ),
                    str(FLOATING_PRICE_PATH.relative_to(PROJECT_ROOT.parent)),
                ],
                "output_files": [],
                "figure_files": [],
                "metrics_summary": run.metrics,
                "warnings": [
                    *(
                        [
                            "future prices are unknown and settlements use actual delivery-interval prices",
                            "shared planning storage trajectory is a tractable proxy for the real-time feedback policy",
                        ]
                        if suite == "extension"
                        else [
                            FORMAL_PRICE_INFORMATION,
                            "正式 Q4 不另建价格预测或价荷联合情景；上述完整价格路径仅用于题给数据下的离线重算。",
                        ]
                    ),
                ],
                "errors": [],
            }
            for run in runs
        ],
        "comparison": {
            "q42_main_minus_baseline_cost_yuan": q42_main.metrics["actual_total_cost_yuan"] - q42_base.metrics["actual_total_cost_yuan"],
            "q42_main_cost_change_vs_baseline_pct": (q42_main.metrics["actual_total_cost_yuan"] / q42_base.metrics["actual_total_cost_yuan"] - 1.0) * 100.0,
            "q43_main_minus_baseline_cost_yuan": q43_main.metrics["actual_total_cost_yuan"] - q43_base.metrics["actual_total_cost_yuan"],
            "q43_main_cost_change_vs_baseline_pct": (q43_main.metrics["actual_total_cost_yuan"] / q43_base.metrics["actual_total_cost_yuan"] - 1.0) * 100.0,
            "q43_minus_q42_main_cost_yuan": q43_main.metrics["actual_total_cost_yuan"] - q42_main.metrics["actual_total_cost_yuan"],
            "q43_cost_change_vs_q42_main_pct": (q43_main.metrics["actual_total_cost_yuan"] / q42_main.metrics["actual_total_cost_yuan"] - 1.0) * 100.0,
        },
        "fallback_trigger": {
            "fallback_id": "F4",
            "condition": "any DA/TR stage failure or hard constraint violation",
            "observed": any(run.metrics["fallback_event_count"] for run in runs),
            "evidence": {run.config.strategy_id: run.validation["fallback_events"] for run in runs if run.validation["fallback_events"]},
        },
        "input_hashes": data.hashes,
        "environment": environment_info(),
        "validation": {run.config.strategy_id: run.validation for run in runs},
        "output_degeneracy": degeneracy,
        "output_files": output_files,
    }
    write_json(round_dir / "run_summary.json", run_summary)


def main() -> None:
    parser = argparse.ArgumentParser(description="Solve Q4 by inheriting Q2/Q3 and replacing only the price matrix")
    parser.add_argument("--suite", choices=("pilot", "core", "extension"), default="core")
    parser.add_argument("--max-days", type=int, default=None)
    args = parser.parse_args()
    run_suite(args.suite, args.max_days)


if __name__ == "__main__":
    raise SystemExit("请从辅助材料根目录运行 run.py，选择对应问题。")
