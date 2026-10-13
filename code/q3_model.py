from __future__ import annotations

import argparse
import hashlib
import json
import math
import platform
import sys
import time as time_module
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]

import numpy as np
import pandas as pd
import scipy
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import csr_matrix, lil_matrix


N_DAYS = 365
N_INTERVALS = 144
RELEASE_HOURS = (0, 6, 12, 18)
RELEASE_STARTS = tuple(hour * 6 for hour in RELEASE_HOURS)
DELTA_HOURS = 1.0 / 6.0
ETA_CHARGE = 0.9
ETA_DISCHARGE = 0.9
SOC_INITIAL_KWH = 6000.0
SOC_MIN_KWH = 1200.0
SOC_MAX_KWH = 10800.0
INTERVAL_LIMIT_KWH = 5000.0 * DELTA_HOURS
EVALUATION_START_INDEX = 31
FINAL_DAY_INDEX = 364
PRIMARY_TOLERANCE = 1e-6
CHECK_TOLERANCE = 5e-5
EVENT_TOLERANCE = 1e-6
RANDOM_SEED = 2026

SOURCE_ROOT = PROJECT_ROOT / "data"
PRICE_PATH = SOURCE_ROOT / "附件1.xlsx"
ACTUAL_PATH = SOURCE_ROOT / "附件2.xlsx"
FORECAST_PATH = SOURCE_ROOT / "附件3.xlsx"
ROUND_DIR = PROJECT_ROOT / "results" / "Q3" / "experiments" / "round4"
TABLE_DIR = ROUND_DIR / "tables"
METRIC_DIR = ROUND_DIR / "metrics"


@dataclass(frozen=True)
class Q3Data:
    dates: list[date]
    price: np.ndarray
    load_kw: np.ndarray
    pv_kw: np.ndarray
    net_load_kwh: np.ndarray
    official_hourly_kw: np.ndarray
    interval_starts: list[str]
    interval_ends: list[str]
    interval_labels: list[str]
    hashes: dict[str, str]


@dataclass(frozen=True)
class ForecastArchive:
    archive_id: str
    load_kw: np.ndarray
    pv_kw: np.ndarray
    net_kwh: np.ndarray
    errors_kwh: np.ndarray
    valid: np.ndarray


@dataclass(frozen=True)
class PolicyConfig:
    strategy_id: str
    role: str
    alpha: float | None = 0.8
    gamma: float = 0.5
    selected_release_hours: tuple[int, ...] = RELEASE_HOURS
    allow_adjustment: bool = True
    fixed_da_storage: bool = False
    feedback_execution: bool = True
    forecast_source: str = "official"
    collect_ledger: bool = False


@dataclass
class LPSolution:
    purchase: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    planned_emergency: np.ndarray
    planned_unused: np.ndarray
    soc_end: np.ndarray
    stage1_objective: float
    stage2_objective: float
    throughput: float
    max_balance_residual: float
    max_storage_residual: float
    max_adjustment_decomposition_residual: float
    primary_tolerance: float
    stage1_status: str
    stage2_status: str


@dataclass
class StrategyRun:
    config: PolicyConfig
    runtime_seconds: float
    metrics: dict[str, Any]
    validation: dict[str, Any]
    daily: pd.DataFrame
    ledger: pd.DataFrame | None
    adjustments: pd.DataFrame | None


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def save_frame(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, index=False, encoding="utf-8-sig", float_format="%.10f")


def _clock(value: Any) -> str:
    if isinstance(value, time):
        return f"{value.hour:02d}:{value.minute:02d}"
    return str(value).strip()


def _fmt_minutes(minutes: int) -> str:
    if minutes == 1440:
        return "24:00"
    hour, minute = divmod(minutes, 60)
    return f"{hour:02d}:{minute:02d}"


def _expected_end_labels(index: int) -> set[str]:
    minutes = (index + 1) * 10
    if minutes == 1440:
        return {"0:00+1", "00:00+1", "24:00"}
    value = _fmt_minutes(minutes)
    return {value, value.lstrip("0")}


def _to_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()


def load_data() -> Q3Data:
    for path in (PRICE_PATH, ACTUAL_PATH, FORECAST_PATH):
        if not path.is_file():
            raise FileNotFoundError(path)

    price_wb = load_workbook(PRICE_PATH, read_only=True, data_only=True)
    price_rows = list(price_wb.active.values)
    if tuple(price_rows[0]) != ("时间", "电价", "小区负载", "光伏发电预测功率"):
        raise ValueError(f"Unexpected attachment 1 header: {price_rows[0]!r}")
    if len(price_rows) != N_INTERVALS + 1:
        raise ValueError("Attachment 1 must contain 144 intervals")
    price = np.asarray([row[1] for row in price_rows[1:]], dtype=float)
    for index, row in enumerate(price_rows[1:]):
        if _clock(row[0]) not in _expected_end_labels(index):
            raise ValueError(f"Attachment 1 time mismatch at interval {index + 1}")
    if not np.isfinite(price).all() or np.any(price <= 0):
        raise ValueError("Prices must be finite and positive")

    actual_wb = load_workbook(ACTUAL_PATH, read_only=True, data_only=True)
    if actual_wb.sheetnames != ["小区负载", "光伏发电实际功率"]:
        raise ValueError(f"Unexpected attachment 2 sheets: {actual_wb.sheetnames}")
    matrices: dict[str, np.ndarray] = {}
    dates: list[date] | None = None
    labels: list[str] | None = None
    for sheet_name in actual_wb.sheetnames:
        rows = list(actual_wb[sheet_name].values)
        if len(rows) != N_DAYS + 1 or len(rows[0]) != N_INTERVALS + 1:
            raise ValueError(f"Unexpected attachment 2 shape in {sheet_name}")
        current_dates = [_to_date(row[0]) for row in rows[1:]]
        current_labels = [_clock(value) for value in rows[0][1:]]
        if dates is None:
            dates, labels = current_dates, current_labels
        elif dates != current_dates or labels != current_labels:
            raise ValueError("Attachment 2 axes differ between sheets")
        matrices[sheet_name] = np.asarray([row[1:] for row in rows[1:]], dtype=float)
    assert dates is not None and labels is not None
    expected_dates = [date(2025, 1, 1) + timedelta(days=index) for index in range(N_DAYS)]
    if dates != expected_dates:
        raise ValueError("Attachment 2 does not cover every day of 2025 in order")
    for index, label in enumerate(labels):
        if label not in _expected_end_labels(index):
            raise ValueError(f"Attachment 2 time mismatch at interval {index + 1}")
    load_kw = matrices["小区负载"]
    pv_kw = matrices["光伏发电实际功率"]
    if not np.isfinite(load_kw).all() or not np.isfinite(pv_kw).all():
        raise ValueError("Attachment 2 contains missing/non-finite values")
    if np.any(load_kw < 0) or np.any(pv_kw < 0):
        raise ValueError("Attachment 2 physical powers must be nonnegative")

    forecast_wb = load_workbook(FORECAST_PATH, read_only=True, data_only=True)
    rows = list(forecast_wb.active.values)
    expected_header = ("日期", "预报时刻", *[f"预报{lead}小时" for lead in range(1, 25)])
    if tuple(rows[0]) != expected_header or len(rows) != N_DAYS * 4 + 1:
        raise ValueError("Unexpected attachment 3 header or row count")
    official = np.full((N_DAYS, 4, 24), np.nan, dtype=float)
    current_date: date | None = None
    seen: set[tuple[date, int]] = set()
    for row in rows[1:]:
        if row[0] not in (None, ""):
            current_date = datetime.strptime(str(row[0]).strip(), "%Y-%m-%d").date()
        if current_date is None:
            raise ValueError("Attachment 3 begins with a blank date")
        release_text = str(row[1]).strip()
        try:
            release_index = [f"{hour}:00" for hour in RELEASE_HOURS].index(release_text)
        except ValueError as exc:
            raise ValueError(f"Unexpected release time {release_text!r}") from exc
        day_index = (current_date - date(2025, 1, 1)).days
        if not 0 <= day_index < N_DAYS or (current_date, release_index) in seen:
            raise ValueError("Attachment 3 has duplicate or out-of-range release")
        values = np.asarray(row[2:], dtype=float)
        if len(values) != 24 or not np.isfinite(values).all() or np.any(values < 0):
            raise ValueError(f"Invalid attachment 3 forecast at {current_date} {release_text}")
        official[day_index, release_index] = values
        seen.add((current_date, release_index))
    if np.isnan(official).any() or len(seen) != N_DAYS * 4:
        raise ValueError("Attachment 3 release coverage is incomplete")

    starts = [_fmt_minutes(index * 10) for index in range(N_INTERVALS)]
    ends = [_fmt_minutes((index + 1) * 10) for index in range(N_INTERVALS)]
    return Q3Data(
        dates=dates,
        price=price,
        load_kw=load_kw,
        pv_kw=pv_kw,
        net_load_kwh=(load_kw - pv_kw) * DELTA_HOURS,
        official_hourly_kw=official,
        interval_starts=starts,
        interval_ends=ends,
        interval_labels=[f"{start}-{end}" for start, end in zip(starts, ends)],
        hashes={"attachment1": sha256(PRICE_PATH), "attachment2": sha256(ACTUAL_PATH), "attachment3": sha256(FORECAST_PATH)},
    )


def build_load_forecast(data: Q3Data) -> tuple[np.ndarray, np.ndarray]:
    predicted = np.full_like(data.load_kw, np.nan)
    valid = np.zeros(N_DAYS, dtype=bool)
    for day_index in range(1, N_DAYS):
        candidates = np.arange(max(0, day_index - 28), day_index, dtype=int)
        target_weekday = data.dates[day_index].weekday()
        same_weekday = np.asarray(
            [idx for idx in candidates if data.dates[idx].weekday() == target_weekday], dtype=int
        )
        if len(same_weekday):
            indices = same_weekday
        else:
            target_weekend = target_weekday >= 5
            same_day_type = np.asarray(
                [idx for idx in candidates if (data.dates[idx].weekday() >= 5) == target_weekend], dtype=int
            )
            indices = (
                same_day_type
                if len(same_day_type)
                else np.arange(max(0, day_index - 7), day_index, dtype=int)
            )
        weights = np.exp(-(day_index - indices) / 7.0)
        predicted[day_index] = np.average(data.load_kw[indices], axis=0, weights=weights)
        valid[day_index] = True
    return predicted, valid


def _interpolate_official_day(data: Q3Data, day_index: int, release_index: int) -> np.ndarray:
    release_hour = RELEASE_HOURS[release_index]
    start = RELEASE_STARTS[release_index]
    if release_hour == 0:
        anchor = 0.0 if day_index == 0 else float(data.pv_kw[day_index - 1, -1])
    else:
        anchor = float(data.pv_kw[day_index, start - 1])
    knots = np.r_[anchor, data.official_hourly_kw[day_index, release_index]]
    targets = (np.arange(start, N_INTERVALS, dtype=float) + 1.0) * DELTA_HOURS - release_hour
    lower = np.floor(targets + 1e-12).astype(int)
    lower = np.minimum(lower, 23)
    fraction = targets - lower
    values = (1.0 - fraction) * knots[lower] + fraction * knots[lower + 1]
    result = np.full(N_INTERVALS, np.nan, dtype=float)
    result[start:] = np.maximum(values, 0.0)
    return result


def build_forecast_archive(data: Q3Data, source: str) -> ForecastArchive:
    load_forecast, valid_days = build_load_forecast(data)
    pv_forecast = np.full((N_DAYS, 4, N_INTERVALS), np.nan, dtype=float)
    if source == "official":
        for day_index in range(N_DAYS):
            for release_index in range(4):
                pv_forecast[day_index, release_index] = _interpolate_official_day(data, day_index, release_index)
    elif source == "historical":
        for day_index in range(1, N_DAYS):
            indices = np.arange(max(0, day_index - 7), day_index, dtype=int)
            weights = np.exp(-(day_index - indices) / 7.0)
            predicted = np.maximum(0.0, np.average(data.pv_kw[indices], axis=0, weights=weights))
            for release_index, start in enumerate(RELEASE_STARTS):
                pv_forecast[day_index, release_index, start:] = predicted[start:]
    else:
        raise ValueError(f"Unknown forecast source {source!r}")

    net = np.full_like(pv_forecast, np.nan)
    errors = np.full_like(pv_forecast, np.nan)
    valid = np.zeros((N_DAYS, 4, N_INTERVALS), dtype=bool)
    for day_index in range(1, N_DAYS):
        for release_index, start in enumerate(RELEASE_STARTS):
            net[day_index, release_index, start:] = (
                load_forecast[day_index, start:] - pv_forecast[day_index, release_index, start:]
            ) * DELTA_HOURS
            errors[day_index, release_index, start:] = data.net_load_kwh[day_index, start:] - net[day_index, release_index, start:]
            valid[day_index, release_index, start:] = valid_days[day_index]
    return ForecastArchive(source, load_forecast, pv_forecast, net, errors, valid)


def build_risk_margins(archive: ForecastArchive, alpha: float | None) -> np.ndarray:
    margins = np.zeros_like(archive.net_kwh)
    if alpha is None:
        return margins
    for day_index in range(1, N_DAYS):
        prior = np.arange(max(1, day_index - 28), day_index, dtype=int)
        if len(prior) < 7:
            continue
        for release_index, start in enumerate(RELEASE_STARTS):
            for interval in range(start, N_INTERVALS):
                eligible = prior[archive.valid[prior, release_index, interval]]
                if len(np.unique(eligible)) < 7:
                    continue
                value = float(np.quantile(archive.errors_kwh[eligible, release_index, interval], alpha, method="linear"))
                margins[day_index, release_index, interval] = max(0.0, value)
    return margins


def incremental_adjustment_fee(
    new_contract: np.ndarray,
    previous_contract: np.ndarray,
    price: np.ndarray,
) -> np.ndarray:
    increase = np.maximum(new_contract - previous_contract, 0.0)
    decrease = np.maximum(previous_contract - new_contract, 0.0)
    return 1.5 * price * increase - 0.5 * price * decrease


class RollingLP:
    def __init__(self, price: np.ndarray) -> None:
        self.price = np.asarray(price, dtype=float)
        self.da_cache: dict[int, tuple[csr_matrix, list[tuple[float | None, float | None]], np.ndarray]] = {}
        self.tr_cache: dict[int, tuple[csr_matrix, np.ndarray]] = {}

    def _da_structure(self, start: int) -> tuple[csr_matrix, list[tuple[float | None, float | None]], np.ndarray]:
        if start in self.da_cache:
            return self.da_cache[start]
        m = N_INTERVALS - start
        a_eq = lil_matrix((2 * m, 5 * m), dtype=float)
        for t in range(m):
            a_eq[t, [t, m + t, 2 * m + t, 3 * m + t]] = [1.0, -1.0, 1.0, -1.0]
            a_eq[m + t, [m + t, 2 * m + t, 4 * m + t]] = [-ETA_CHARGE, 1.0 / ETA_DISCHARGE, 1.0]
            if t:
                a_eq[m + t, 4 * m + t - 1] = -1.0
        bounds = (
            [(0.0, None)] * m
            + [(0.0, INTERVAL_LIMIT_KWH)] * m
            + [(0.0, INTERVAL_LIMIT_KWH)] * m
            + [(0.0, None)] * m
            + [(SOC_MIN_KWH, SOC_MAX_KWH)] * m
        )
        secondary = np.r_[np.zeros(m), np.ones(m), np.ones(m), np.zeros(2 * m)]
        value = (a_eq.tocsr(), bounds, secondary)
        self.da_cache[start] = value
        return value

    def solve_da(self, planned_net: np.ndarray, initial_soc: float, terminal_value: float) -> LPSolution:
        start = 0
        m = N_INTERVALS
        a_eq, bounds, secondary = self._da_structure(start)
        b_eq = np.r_[planned_net, np.zeros(m)]
        b_eq[m] = initial_soc
        primary = np.r_[self.price, np.zeros(4 * m)]
        primary[-1] = -terminal_value
        stage1 = linprog(primary, A_eq=a_eq, b_eq=b_eq, bounds=bounds, method="highs")
        if not stage1.success:
            raise RuntimeError(f"DA stage1 failed: {stage1.message}")
        stage2 = linprog(
            secondary,
            A_ub=csr_matrix(primary.reshape(1, -1)),
            b_ub=np.asarray([float(stage1.fun) + PRIMARY_TOLERANCE]),
            A_eq=a_eq,
            b_eq=b_eq,
            bounds=bounds,
            method="highs",
        )
        if not stage2.success:
            raise RuntimeError(f"DA stage2 failed: {stage2.message}")
        q, charge, discharge, unused, soc = stage2.x.reshape(5, m)
        balance = q - charge + discharge - unused - planned_net
        prior_soc = np.r_[initial_soc, soc[:-1]]
        storage = soc - prior_soc - ETA_CHARGE * charge + discharge / ETA_DISCHARGE
        return LPSolution(
            q,
            charge,
            discharge,
            np.zeros(m),
            unused,
            soc,
            float(stage1.fun),
            float(primary @ stage2.x),
            float(charge.sum() + discharge.sum()),
            float(np.max(np.abs(balance))),
            float(np.max(np.abs(storage))),
            0.0,
            PRIMARY_TOLERANCE,
            str(stage1.message),
            str(stage2.message),
        )

    def _tr_structure(self, start: int) -> tuple[csr_matrix, np.ndarray]:
        if start in self.tr_cache:
            return self.tr_cache[start]
        m = N_INTERVALS - start
        # Variables by m-block: y, increase, decrease, charge, discharge,
        # planned emergency, planned unused, SOC.
        a_eq = lil_matrix((3 * m, 8 * m), dtype=float)
        for t in range(m):
            a_eq[t, [t, 3 * m + t, 4 * m + t, 5 * m + t, 6 * m + t]] = [1.0, -1.0, 1.0, 1.0, -1.0]
            a_eq[m + t, [3 * m + t, 4 * m + t, 7 * m + t]] = [-ETA_CHARGE, 1.0 / ETA_DISCHARGE, 1.0]
            if t:
                a_eq[m + t, 7 * m + t - 1] = -1.0
            a_eq[2 * m + t, [t, m + t, 2 * m + t]] = [1.0, -1.0, 1.0]
        secondary = np.r_[np.zeros(3 * m), np.ones(m), np.ones(m), np.zeros(3 * m)]
        value = (a_eq.tocsr(), secondary)
        self.tr_cache[start] = value
        return value

    def solve_tr(
        self,
        start: int,
        planned_net: np.ndarray,
        initial_soc: float,
        previous_contract: np.ndarray,
        terminal_value: float,
        allow_adjustment: bool,
        fixed_charge: np.ndarray | None = None,
        fixed_discharge: np.ndarray | None = None,
    ) -> LPSolution:
        m = N_INTERVALS - start
        a_eq, secondary = self._tr_structure(start)
        local_price = self.price[start:]
        reference = previous_contract[start:]
        b_eq = np.r_[planned_net, np.zeros(m), reference]
        b_eq[m] = initial_soc
        primary = np.r_[
            np.zeros(m),
            1.5 * local_price,
            -0.5 * local_price,
            np.zeros(2 * m),
            5.0 * local_price,
            np.zeros(2 * m),
        ]
        primary[-1] = -terminal_value
        y_bounds = [(0.0, None)] * m if allow_adjustment else [(float(value), float(value)) for value in reference]
        charge_bounds = (
            [(0.0, INTERVAL_LIMIT_KWH)] * m
            if fixed_charge is None
            else [(float(value), float(value)) for value in fixed_charge]
        )
        discharge_bounds = (
            [(0.0, INTERVAL_LIMIT_KWH)] * m
            if fixed_discharge is None
            else [(float(value), float(value)) for value in fixed_discharge]
        )
        bounds = (
            y_bounds
            + [(0.0, None)] * (2 * m)
            + charge_bounds
            + discharge_bounds
            + [(0.0, None)] * m
            + [(0.0, None)] * m
            + [(SOC_MIN_KWH, SOC_MAX_KWH)] * m
        )
        stage1 = linprog(
            primary,
            A_eq=a_eq,
            b_eq=b_eq,
            bounds=bounds,
            method="highs",
        )
        if not stage1.success:
            raise RuntimeError(f"TR {RELEASE_HOURS[RELEASE_STARTS.index(start)]}:00 stage1 failed: {stage1.message}")
        stage2 = linprog(
            secondary,
            A_ub=csr_matrix(primary.reshape(1, -1)),
            b_ub=np.asarray([float(stage1.fun) + PRIMARY_TOLERANCE]),
            A_eq=a_eq,
            b_eq=b_eq,
            bounds=bounds,
            method="highs",
        )
        if not stage2.success:
            raise RuntimeError(f"TR {RELEASE_HOURS[RELEASE_STARTS.index(start)]}:00 stage2 failed: {stage2.message}")
        y, increase, decrease, charge, discharge, emergency, unused, soc = stage2.x.reshape(8, m)
        balance = y - charge + discharge + emergency - unused - planned_net
        prior_soc = np.r_[initial_soc, soc[:-1]]
        storage = soc - prior_soc - ETA_CHARGE * charge + discharge / ETA_DISCHARGE
        adjustment_decomposition = y - increase + decrease - reference
        adjustment_overlap = np.minimum(increase, decrease)
        return LPSolution(
            y,
            charge,
            discharge,
            emergency,
            unused,
            soc,
            float(stage1.fun),
            float(primary @ stage2.x),
            float(charge.sum() + discharge.sum()),
            float(np.max(np.abs(balance))),
            float(np.max(np.abs(storage))),
            float(max(np.max(np.abs(adjustment_decomposition)), np.max(adjustment_overlap))),
            PRIMARY_TOLERANCE,
            str(stage1.message),
            str(stage2.message),
        )


def _fallback_da(planned_net: np.ndarray, initial_soc: float) -> LPSolution:
    q = np.maximum(planned_net, 0.0)
    zeros = np.zeros(N_INTERVALS)
    return LPSolution(
        q,
        zeros.copy(),
        zeros.copy(),
        zeros.copy(),
        np.maximum(-planned_net, 0.0),
        np.full(N_INTERVALS, initial_soc),
        math.nan,
        math.nan,
        0.0,
        0.0,
        0.0,
        0.0,
        PRIMARY_TOLERANCE,
        "failed_fallback_activated",
        "failed_fallback_activated",
    )


def run_policy(
    data: Q3Data,
    archive: ForecastArchive,
    margins: np.ndarray,
    config: PolicyConfig,
) -> StrategyRun:
    started = time_module.perf_counter()
    solver = RollingLP(data.price)
    selected = set(config.selected_release_hours)
    daily_rows: list[dict[str, Any]] = []
    ledger_rows: list[dict[str, Any]] = []
    adjustment_event_rows: list[dict[str, Any]] = []
    current_soc = SOC_INITIAL_KWH
    max_checks: dict[str, float] = {}
    fallback_events: list[dict[str, Any]] = []
    plan_count = 0

    def update_max(name: str, value: float) -> None:
        if math.isfinite(value):
            max_checks[name] = max(max_checks.get(name, 0.0), float(value))

    for day_index, day in enumerate(data.dates):
        day_initial_soc = current_soc
        q = np.zeros(N_INTERVALS)
        final_y = np.zeros(N_INTERVALS)
        exec_charge = np.zeros(N_INTERVALS)
        exec_discharge = np.zeros(N_INTERVALS)
        exec_plan_emergency = np.zeros(N_INTERVALS)
        exec_plan_unused = np.zeros(N_INTERVALS)
        actual_emergency = np.zeros(N_INTERVALS)
        actual_unused = np.zeros(N_INTERVALS)
        adjustment_fee_by_interval = np.zeros(N_INTERVALS)
        adjustment_increase_by_interval = np.zeros(N_INTERVALS)
        adjustment_decrease_by_interval = np.zeros(N_INTERVALS)
        adjustment_count_by_interval = np.zeros(N_INTERVALS, dtype=int)
        soc_end = np.zeros(N_INTERVALS)
        plan_release = np.zeros(N_INTERVALS, dtype=int)
        terminal_value = 0.0 if day_index == FINAL_DAY_INDEX else config.gamma * ETA_DISCHARGE * float(np.median(data.price))

        if day_index == 0:
            da = _fallback_da(np.zeros(N_INTERVALS), current_soc)
            da.stage1_status = "cold_start_no_load_history"
            da.stage2_status = "cold_start_no_load_history"
        else:
            planned_da = archive.net_kwh[day_index, 0] + margins[day_index, 0]
            try:
                da = solver.solve_da(planned_da, current_soc, terminal_value)
                plan_count += 1
            except RuntimeError as exc:
                da = _fallback_da(planned_da, current_soc)
                fallback_events.append({"date": day.isoformat(), "release_hour": 0, "error": str(exc)})
        q[:] = da.purchase
        final_y[:] = q
        current_charge = da.charge.copy()
        current_discharge = da.discharge.copy()
        current_plan_emergency = da.planned_emergency.copy()
        current_plan_unused = da.planned_unused.copy()
        da_charge = da.charge.copy()
        da_discharge = da.discharge.copy()
        update_max("planning_balance_max_abs", da.max_balance_residual)
        update_max("storage_recurrence_max_abs", da.max_storage_residual)
        update_max("stage2_objective_excess", max(0.0, da.stage2_objective - da.stage1_objective) if day_index else 0.0)

        for release_index, (release_hour, start) in enumerate(zip(RELEASE_HOURS, RELEASE_STARTS)):
            if day_index > 0 and release_hour > 0 and release_hour in selected:
                planned = archive.net_kwh[day_index, release_index, start:] + margins[day_index, release_index, start:]
                try:
                    tr = solver.solve_tr(
                        start,
                        planned,
                        current_soc,
                        final_y,
                        terminal_value,
                        config.allow_adjustment,
                        da_charge[start:] if config.fixed_da_storage else None,
                        da_discharge[start:] if config.fixed_da_storage else None,
                    )
                    plan_count += 1
                    previous_y = final_y[start:].copy()
                    new_y = tr.purchase
                    increase = np.maximum(new_y - previous_y, 0.0)
                    decrease = np.maximum(previous_y - new_y, 0.0)
                    incremental_fee = incremental_adjustment_fee(new_y, previous_y, data.price[start:])
                    changed = np.abs(new_y - previous_y) > EVENT_TOLERANCE
                    adjustment_fee_by_interval[start:] += incremental_fee
                    adjustment_increase_by_interval[start:] += increase
                    adjustment_decrease_by_interval[start:] += decrease
                    adjustment_count_by_interval[start:] += changed.astype(int)
                    if config.collect_ledger:
                        for local_index, interval in enumerate(range(start, N_INTERVALS)):
                            adjustment_event_rows.append(
                                {
                                    "strategy_id": config.strategy_id,
                                    "date": day.isoformat(),
                                    "release_hour": release_hour,
                                    "interval_index": interval + 1,
                                    "interval_label": data.interval_labels[interval],
                                    "price_yuan_per_kwh": data.price[interval],
                                    "previous_contract_kwh": previous_y[local_index],
                                    "new_contract_kwh": new_y[local_index],
                                    "increase_kwh": increase[local_index],
                                    "decrease_kwh": decrease[local_index],
                                    "incremental_adjustment_fee_yuan": incremental_fee[local_index],
                                    "changed": bool(changed[local_index]),
                                }
                            )
                    final_y[start:] = new_y
                    current_charge[start:] = tr.charge
                    current_discharge[start:] = tr.discharge
                    current_plan_emergency[start:] = tr.planned_emergency
                    current_plan_unused[start:] = tr.planned_unused
                    plan_release[start:] = release_hour
                    update_max("planning_balance_max_abs", tr.max_balance_residual)
                    update_max("storage_recurrence_max_abs", tr.max_storage_residual)
                    update_max("adjustment_decomposition_violation", tr.max_adjustment_decomposition_residual)
                    update_max("stage2_objective_excess", max(0.0, tr.stage2_objective - tr.stage1_objective))
                    update_max("planned_emergency_main_max", float(np.max(tr.planned_emergency)) if config.allow_adjustment else 0.0)
                except RuntimeError as exc:
                    fallback_events.append({"date": day.isoformat(), "release_hour": release_hour, "error": str(exc)})

            end = RELEASE_STARTS[release_index + 1] if release_index < 3 else N_INTERVALS
            for interval in range(start, end):
                # The contract is frozen between forecast releases, but the problem
                # does not freeze storage.  From Jan 2 onward, storage reacts only
                # to the current interval's observed net load and current SOC.  No
                # future actual value is visible to this rule.
                net_gap = data.net_load_kwh[day_index, interval] - final_y[interval]
                if day_index > 0 and config.feedback_execution:
                    charge = min(
                        max(-net_gap, 0.0),
                        INTERVAL_LIMIT_KWH,
                        max(0.0, (SOC_MAX_KWH - current_soc) / ETA_CHARGE),
                    )
                    discharge = min(
                        max(net_gap, 0.0),
                        INTERVAL_LIMIT_KWH,
                        max(0.0, ETA_DISCHARGE * (current_soc - SOC_MIN_KWH)),
                    )
                else:
                    charge = current_charge[interval]
                    discharge = current_discharge[interval]
                residual = net_gap + charge - discharge
                emergency = max(residual, 0.0)
                unused = max(-residual, 0.0)
                next_soc = current_soc + ETA_CHARGE * charge - discharge / ETA_DISCHARGE
                exec_charge[interval] = charge
                exec_discharge[interval] = discharge
                exec_plan_emergency[interval] = current_plan_emergency[interval]
                exec_plan_unused[interval] = current_plan_unused[interval]
                actual_emergency[interval] = emergency
                actual_unused[interval] = unused
                soc_end[interval] = next_soc
                update_max("soc_lower_violation", max(0.0, SOC_MIN_KWH - next_soc))
                update_max("soc_upper_violation", max(0.0, next_soc - SOC_MAX_KWH))
                update_max("charge_limit_violation", max(0.0, charge - INTERVAL_LIMIT_KWH))
                update_max("discharge_limit_violation", max(0.0, discharge - INTERVAL_LIMIT_KWH))
                update_max("simultaneous_charge_discharge_max", min(charge, discharge))
                update_max("actual_emergency_unused_overlap_max", min(emergency, unused))
                balance = (
                    final_y[interval]
                    + emergency
                    + data.pv_kw[day_index, interval] * DELTA_HOURS
                    + discharge
                    - data.load_kw[day_index, interval] * DELTA_HOURS
                    - charge
                    - unused
                )
                update_max("actual_balance_max_abs", abs(balance))
                recurrence = next_soc - current_soc - ETA_CHARGE * charge + discharge / ETA_DISCHARGE
                update_max("actual_storage_recurrence_max_abs", abs(recurrence))
                current_soc = next_soc

        original_cost_by_interval = data.price * q
        original_cost = float(original_cost_by_interval.sum())
        contract_fee_by_interval = original_cost_by_interval + adjustment_fee_by_interval
        adjustment_net = float(adjustment_fee_by_interval.sum())
        contract_cost = float(contract_fee_by_interval.sum())
        emergency_cost = float(5.0 * data.price @ actual_emergency)
        total_cost = contract_cost + emergency_cost
        daily_rows.append(
            {
                "strategy_id": config.strategy_id,
                "date": day.isoformat(),
                "initial_soc_kwh": day_initial_soc,
                "terminal_soc_kwh": current_soc,
                "original_contract_kwh": float(q.sum()),
                "final_contract_kwh": float(final_y.sum()),
                "adjustment_increase_kwh": float(adjustment_increase_by_interval.sum()),
                "adjustment_decrease_kwh": float(adjustment_decrease_by_interval.sum()),
                "adjustment_transaction_count": int(adjustment_count_by_interval.sum()),
                "actual_total_purchase_kwh": float(final_y.sum() + actual_emergency.sum()),
                "charge_kwh": float(exec_charge.sum()),
                "discharge_kwh": float(exec_discharge.sum()),
                "planned_emergency_kwh": float(exec_plan_emergency.sum()),
                "planned_unused_kwh": float(exec_plan_unused.sum()),
                "emergency_purchase_kwh": float(actual_emergency.sum()),
                "actual_unused_kwh": float(actual_unused.sum()),
                "original_contract_cost_yuan": original_cost,
                "adjustment_net_cost_yuan": adjustment_net,
                "contract_total_cost_yuan": contract_cost,
                "emergency_cost_yuan": emergency_cost,
                "actual_total_cost_yuan": total_cost,
                "risk_margin_kwh": float(
                    sum(
                        margins[
                            day_index,
                            RELEASE_HOURS.index(int(plan_release[interval])),
                            interval,
                        ]
                        for interval in range(N_INTERVALS)
                    )
                ),
                "fallback_used": any(item["date"] == day.isoformat() for item in fallback_events),
            }
        )

        if config.collect_ledger:
            prior_soc = np.r_[day_initial_soc, soc_end[:-1]]
            for interval in range(N_INTERVALS):
                executing_release_index = RELEASE_HOURS.index(int(plan_release[interval]))
                ledger_rows.append(
                    {
                        "strategy_id": config.strategy_id,
                        "date": day.isoformat(),
                        "interval_index": interval + 1,
                        "interval_start": data.interval_starts[interval],
                        "interval_end": data.interval_ends[interval],
                        "interval_label": data.interval_labels[interval],
                        "executing_release_hour": int(plan_release[interval]),
                        "price_yuan_per_kwh": data.price[interval],
                        "actual_load_kw": data.load_kw[day_index, interval],
                        "actual_pv_kw": data.pv_kw[day_index, interval],
                        "actual_net_load_kwh": data.net_load_kwh[day_index, interval],
                        "forecast_load_kw": archive.load_kw[day_index, interval] if day_index else np.nan,
                        "forecast_pv_kw": (
                            archive.pv_kw[day_index, executing_release_index, interval]
                            if day_index
                            else np.nan
                        ),
                        "forecast_net_load_kwh": (
                            archive.net_kwh[day_index, executing_release_index, interval]
                            if day_index
                            else np.nan
                        ),
                        "risk_margin_kwh": (
                            margins[day_index, executing_release_index, interval]
                            if day_index
                            else 0.0
                        ),
                        "original_contract_kwh": q[interval],
                        "final_contract_kwh": final_y[interval],
                        "adjustment_increase_kwh": adjustment_increase_by_interval[interval],
                        "adjustment_decrease_kwh": adjustment_decrease_by_interval[interval],
                        "adjustment_transaction_count": int(adjustment_count_by_interval[interval]),
                        "actual_charge_kwh": exec_charge[interval],
                        "actual_discharge_kwh": exec_discharge[interval],
                        "planned_emergency_kwh": exec_plan_emergency[interval],
                        "planned_unused_kwh": exec_plan_unused[interval],
                        "actual_emergency_kwh": actual_emergency[interval],
                        "actual_unused_kwh": actual_unused[interval],
                        "soc_start_kwh": prior_soc[interval],
                        "soc_end_kwh": soc_end[interval],
                        "original_contract_cost_yuan": original_cost_by_interval[interval],
                        "cumulative_adjustment_fee_yuan": adjustment_fee_by_interval[interval],
                        "contract_total_cost_yuan": contract_fee_by_interval[interval],
                        "adjustment_net_cost_yuan": adjustment_fee_by_interval[interval],
                        "emergency_cost_yuan": 5.0 * data.price[interval] * actual_emergency[interval],
                        "actual_total_cost_yuan": contract_fee_by_interval[interval] + 5.0 * data.price[interval] * actual_emergency[interval],
                    }
                )

    daily = pd.DataFrame(daily_rows)
    evaluation = daily.iloc[EVALUATION_START_INDEX:]
    metrics = {
        "strategy_id": config.strategy_id,
        "role": config.role,
        "evaluation_start": data.dates[EVALUATION_START_INDEX].isoformat(),
        "evaluation_end": data.dates[-1].isoformat(),
        "evaluation_days": len(evaluation),
        "alpha": config.alpha,
        "gamma": config.gamma,
        "selected_release_hours": list(config.selected_release_hours),
        "allow_adjustment": config.allow_adjustment,
        "fixed_da_storage": config.fixed_da_storage,
        "feedback_execution": config.feedback_execution,
        "execution_mode": "current_interval_actual_feedback" if config.feedback_execution else "fixed_planned_dispatch",
        "forecast_source": config.forecast_source,
        "initial_soc_2025_02_01_kwh": float(daily.iloc[EVALUATION_START_INDEX]["initial_soc_kwh"]),
        "terminal_soc_2025_12_31_kwh": float(evaluation.iloc[-1]["terminal_soc_kwh"]),
        "original_contract_kwh": float(evaluation["original_contract_kwh"].sum()),
        "final_contract_kwh": float(evaluation["final_contract_kwh"].sum()),
        "adjustment_increase_kwh": float(evaluation["adjustment_increase_kwh"].sum()),
        "adjustment_decrease_kwh": float(evaluation["adjustment_decrease_kwh"].sum()),
        "adjustment_transaction_count": int(evaluation["adjustment_transaction_count"].sum()),
        "emergency_purchase_kwh": float(evaluation["emergency_purchase_kwh"].sum()),
        "actual_total_purchase_kwh": float(evaluation["actual_total_purchase_kwh"].sum()),
        "charge_kwh": float(evaluation["charge_kwh"].sum()),
        "discharge_kwh": float(evaluation["discharge_kwh"].sum()),
        "original_contract_cost_yuan": float(evaluation["original_contract_cost_yuan"].sum()),
        "adjustment_net_cost_yuan": float(evaluation["adjustment_net_cost_yuan"].sum()),
        "contract_total_cost_yuan": float(evaluation["contract_total_cost_yuan"].sum()),
        "emergency_cost_yuan": float(evaluation["emergency_cost_yuan"].sum()),
        "actual_total_cost_yuan": float(evaluation["actual_total_cost_yuan"].sum()),
        "risk_margin_kwh": float(evaluation["risk_margin_kwh"].sum()),
        "fallback_event_count": len(fallback_events),
        "planning_calls": plan_count,
    }
    fee_reconciliation = abs(metrics["contract_total_cost_yuan"] - metrics["original_contract_cost_yuan"] - metrics["adjustment_net_cost_yuan"])
    total_reconciliation = abs(metrics["actual_total_cost_yuan"] - metrics["contract_total_cost_yuan"] - metrics["emergency_cost_yuan"])
    cross_day = np.abs(daily["initial_soc_kwh"].to_numpy()[1:] - daily["terminal_soc_kwh"].to_numpy()[:-1])
    validation = {
        **max_checks,
        "fee_reconciliation_abs_yuan": float(fee_reconciliation),
        "total_cost_reconciliation_abs_yuan": float(total_reconciliation),
        "cross_day_soc_max_abs_kwh": float(np.max(cross_day)),
        "incremental_example_100_to_80": float(
            100.0 + incremental_adjustment_fee(np.array([80.0]), np.array([100.0]), np.array([1.0]))[0]
        ),
        "incremental_example_100_to_80_to_100": float(
            100.0
            + incremental_adjustment_fee(np.array([80.0]), np.array([100.0]), np.array([1.0]))[0]
            + incremental_adjustment_fee(np.array([100.0]), np.array([80.0]), np.array([1.0]))[0]
        ),
        "incremental_example_100_to_120": float(
            100.0 + incremental_adjustment_fee(np.array([120.0]), np.array([100.0]), np.array([1.0]))[0]
        ),
        "fallback_events": fallback_events,
    }
    hard_values = [
        validation.get("planning_balance_max_abs", 0.0),
        validation.get("storage_recurrence_max_abs", 0.0),
        validation.get("actual_balance_max_abs", 0.0),
        validation.get("actual_storage_recurrence_max_abs", 0.0),
        validation.get("soc_lower_violation", 0.0),
        validation.get("soc_upper_violation", 0.0),
        validation.get("charge_limit_violation", 0.0),
        validation.get("discharge_limit_violation", 0.0),
        validation.get("simultaneous_charge_discharge_max", 0.0),
        validation.get("adjustment_decomposition_violation", 0.0),
        validation["fee_reconciliation_abs_yuan"],
        validation["total_cost_reconciliation_abs_yuan"],
        validation["cross_day_soc_max_abs_kwh"],
    ]
    validation["hard_checks_pass"] = bool(max(hard_values) <= CHECK_TOLERANCE and len(fallback_events) == 0)
    return StrategyRun(
        config=config,
        runtime_seconds=time_module.perf_counter() - started,
        metrics=metrics,
        validation=validation,
        daily=daily,
        ledger=pd.DataFrame(ledger_rows) if config.collect_ledger else None,
        adjustments=pd.DataFrame(adjustment_event_rows) if config.collect_ledger else None,
    )


def emergency_events(ledger: pd.DataFrame) -> pd.DataFrame:
    events: list[dict[str, Any]] = []
    for day, group in ledger.groupby("date", sort=True):
        ordered = group.sort_values("interval_index").reset_index(drop=True)
        amounts = ordered["actual_emergency_kwh"].to_numpy(dtype=float)
        active = amounts > EVENT_TOLERANCE
        cursor = 0
        event_index = 0
        while cursor < N_INTERVALS:
            if not active[cursor]:
                cursor += 1
                continue
            start = cursor
            while cursor + 1 < N_INTERVALS and active[cursor + 1]:
                cursor += 1
            end = cursor
            event_index += 1
            events.append(
                {
                    "date": day,
                    "event_index": event_index,
                    "start_interval_index": start + 1,
                    "end_interval_index": end + 1,
                    "emergency_time_range": f"{ordered.loc[start, 'interval_start']}-{ordered.loc[end, 'interval_end']}",
                    "emergency_purchase_kwh": float(amounts[start : end + 1].sum()),
                }
            )
            cursor += 1
    return pd.DataFrame(events)


def storage_summary(ledger: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for day, group in ledger.groupby("date", sort=True):
        ordered = group.sort_values("interval_index").reset_index(drop=True)
        initial_soc = float(ordered.iloc[0]["soc_start_kwh"])
        terminal_soc = float(ordered.iloc[-1]["soc_end_kwh"])
        for block in range(6):
            section = ordered.iloc[block * 24 : (block + 1) * 24]
            rows.append(
                {
                    "date": day,
                    "block_index": block + 1,
                    "time_range": f"{block * 4:02d}:00-{(block + 1) * 4:02d}:00" if block < 5 else "20:00-24:00",
                    "charge_kwh": float(section["actual_charge_kwh"].sum()),
                    "discharge_kwh": float(section["actual_discharge_kwh"].sum()),
                    "soc_00_kwh": initial_soc if block == 0 else np.nan,
                    "soc_24_kwh": terminal_soc if block == 1 else np.nan,
                }
            )
    return pd.DataFrame(rows)


def specified_tables(run: StrategyRun, events: pd.DataFrame, storage: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    if run.ledger is None:
        raise ValueError("Specified tables require a collected ledger")
    specified = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]
    indices = [61, 73, 85, 97, 109, 121]
    rows: list[dict[str, Any]] = []
    for day in specified:
        detail = run.ledger[run.ledger["date"] == day].set_index("interval_index")
        daily = run.daily[run.daily["date"] == day].iloc[0]
        row: dict[str, Any] = {"date": day}
        for index in indices:
            row[f"{detail.loc[index, 'interval_label']}_final_contract_kwh"] = float(detail.loc[index, "final_contract_kwh"])
        row.update(
            {
                "daily_original_contract_kwh": float(daily["original_contract_kwh"]),
                "daily_final_contract_kwh": float(daily["final_contract_kwh"]),
                "daily_emergency_purchase_kwh": float(daily["emergency_purchase_kwh"]),
                "daily_actual_total_purchase_kwh": float(daily["actual_total_purchase_kwh"]),
                "daily_original_contract_cost_yuan": float(daily["original_contract_cost_yuan"]),
                "daily_adjustment_net_cost_yuan": float(daily["adjustment_net_cost_yuan"]),
                "daily_contract_total_cost_yuan": float(daily["contract_total_cost_yuan"]),
                "daily_emergency_cost_yuan": float(daily["emergency_cost_yuan"]),
                "daily_actual_total_cost_yuan": float(daily["actual_total_cost_yuan"]),
            }
        )
        rows.append(row)
    return (
        pd.DataFrame(rows),
        storage[storage["date"].isin(specified)].copy(),
        events[events["date"].isin(specified)].copy(),
    )


def _comparison_frame(runs: list[StrategyRun]) -> pd.DataFrame:
    rows = []
    for run in runs:
        row = {"strategy_id": run.config.strategy_id, "role": run.config.role, **run.metrics}
        row["runtime_seconds"] = run.runtime_seconds
        row["hard_checks_pass"] = run.validation["hard_checks_pass"]
        rows.append(row)
    frame = pd.DataFrame(rows)
    reference = float(frame.iloc[0]["actual_total_cost_yuan"])
    frame["cost_difference_vs_first_yuan"] = frame["actual_total_cost_yuan"] - reference
    frame["cost_change_vs_first_pct"] = frame["cost_difference_vs_first_yuan"] / reference * 100.0
    return frame


def forecast_diagnostics(data: Q3Data, archive: ForecastArchive) -> dict[str, Any]:
    eval_errors = archive.errors_kwh[EVALUATION_START_INDEX:]
    values = eval_errors[np.isfinite(eval_errors)]
    release_metrics = {}
    for index, hour in enumerate(RELEASE_HOURS):
        current = eval_errors[:, index]
        current = current[np.isfinite(current)]
        release_metrics[str(hour)] = {
            "n": int(len(current)),
            "mae_kwh": float(np.mean(np.abs(current))),
            "rmse_kwh": float(np.sqrt(np.mean(current**2))),
            "bias_actual_minus_forecast_kwh": float(np.mean(current)),
            "q80_kwh": float(np.quantile(current, 0.8, method="linear")),
        }
    integer_hour_checks = []
    for day_index in (31, 171, 364):
        for release_index, start in enumerate(RELEASE_STARTS):
            for lead in (1, 6, 12, 24):
                target = start + lead * 6 - 1
                if target < N_INTERVALS:
                    integer_hour_checks.append(
                        abs(archive.pv_kw[day_index, release_index, target] - data.official_hourly_kw[day_index, release_index, lead - 1])
                    )
    return {
        "archive_id": archive.archive_id,
        "all_release_error_n": int(len(values)),
        "all_release_mae_kwh": float(np.mean(np.abs(values))),
        "release_metrics": release_metrics,
        "official_integer_hour_mapping_max_abs_kw": float(max(integer_hour_checks)) if archive.archive_id == "official" else None,
    }


def environment_info() -> dict[str, Any]:
    return {
        "python": sys.version,
        "platform": platform.platform(),
        "numpy": np.__version__,
        "pandas": pd.__version__,
        "scipy": scipy.__version__,
        "solver": "scipy.optimize.linprog(method='highs')",
        "random_seed": RANDOM_SEED,
    }


def run_suite(suite: str) -> None:
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    METRIC_DIR.mkdir(parents=True, exist_ok=True)
    data = load_data()
    official = build_forecast_archive(data, "official")
    historical = build_forecast_archive(data, "historical")
    margin_cache: dict[tuple[str, float | None], np.ndarray] = {}

    def get_margins(archive: ForecastArchive, alpha: float | None) -> np.ndarray:
        key = (archive.archive_id, alpha)
        if key not in margin_cache:
            margin_cache[key] = build_risk_margins(archive, alpha)
        return margin_cache[key]

    core_configs = [
        PolicyConfig("M1I_official_S3_alpha0.8_gamma0.5", "main", collect_ledger=True),
        PolicyConfig("B1I_official_S3_b0_gamma0.5", "usable_baseline", alpha=None),
        PolicyConfig("D_I_frequency_S0", "diagnostic_frequency", selected_release_hours=(0,)),
        PolicyConfig("D_I_frequency_S1", "diagnostic_frequency", selected_release_hours=(0, 6)),
        PolicyConfig("D_I_frequency_S2", "diagnostic_frequency", selected_release_hours=(0, 6, 12)),
        PolicyConfig("D_I_no_adjustment", "diagnostic_mechanism", allow_adjustment=False),
        PolicyConfig("D_I_historical_pv", "diagnostic_mechanism", forecast_source="historical"),
        PolicyConfig("D_I_fixed_plan_execution", "diagnostic_mechanism", feedback_execution=False),
    ]
    sensitivity_configs = [
        PolicyConfig("S_I_alpha0.5", "sensitivity", alpha=0.5),
        PolicyConfig("S_I_alpha0.95", "sensitivity", alpha=0.95),
        PolicyConfig("S_I_gamma0", "sensitivity", gamma=0.0),
        PolicyConfig("S_I_gamma1", "sensitivity", gamma=1.0),
    ]
    configs = sensitivity_configs if suite == "sensitivity" else core_configs + (sensitivity_configs if suite == "full" else [])

    runs: list[StrategyRun] = []
    for config in configs:
        archive = official if config.forecast_source == "official" else historical
        print(f"START {config.strategy_id}", flush=True)
        run = run_policy(data, archive, get_margins(archive, config.alpha), config)
        print(
            f"DONE {config.strategy_id} cost={run.metrics['actual_total_cost_yuan']:.2f} "
            f"emergency={run.metrics['emergency_purchase_kwh']:.3f} runtime={run.runtime_seconds:.2f}s",
            flush=True,
        )
        if not run.validation["hard_checks_pass"]:
            raise RuntimeError(f"Hard checks failed for {config.strategy_id}: {run.validation}")
        runs.append(run)

    if suite == "sensitivity":
        comparison = _comparison_frame(runs)
        save_frame(comparison, TABLE_DIR / "sensitivity_comparison.csv")
        write_json(METRIC_DIR / "sensitivity_validation.json", {run.config.strategy_id: run.validation for run in runs})
        summary_path = ROUND_DIR / "run_summary.json"
        if summary_path.is_file():
            prior = json.loads(summary_path.read_text(encoding="utf-8"))
            for run in runs:
                prior["methods"].append(
                    {
                        "method_id": run.config.strategy_id,
                        "role": run.config.role,
                        "script": "code/Q3/q3_model.py",
                        "status": "success",
                        "execution_time_seconds": run.runtime_seconds,
                        "input_files": [],
                        "output_files": ["tables/sensitivity_comparison.csv"],
                        "figure_files": [],
                        "metrics_summary": run.metrics,
                        "warnings": [],
                        "errors": [],
                    }
                )
                prior["validation"][run.config.strategy_id] = run.validation
            prior["suite"] = "core_plus_sensitivity"
            prior["output_files"].extend(["tables/sensitivity_comparison.csv", "metrics/sensitivity_validation.json"])
            write_json(summary_path, prior)
        return

    main_run = runs[0]
    baseline_run = runs[1]
    assert main_run.ledger is not None
    assert main_run.adjustments is not None
    evaluation_ledger = main_run.ledger[main_run.ledger["date"] >= "2025-02-01"].copy()
    evaluation_adjustments = main_run.adjustments[main_run.adjustments["date"] >= "2025-02-01"].copy()
    evaluation_daily = main_run.daily[main_run.daily["date"] >= "2025-02-01"].copy()
    events = emergency_events(evaluation_ledger)
    storage = storage_summary(evaluation_ledger)
    table1, table2, table3 = specified_tables(main_run, events, storage)

    save_frame(evaluation_ledger, TABLE_DIR / "q3_main_interval_ledger.csv")
    save_frame(evaluation_adjustments, TABLE_DIR / "q3_adjustment_events.csv")
    save_frame(evaluation_daily, TABLE_DIR / "q3_main_daily_summary.csv")
    save_frame(baseline_run.daily[baseline_run.daily["date"] >= "2025-02-01"], TABLE_DIR / "q3_baseline_daily_summary.csv")
    save_frame(events, TABLE_DIR / "q3_emergency_events.csv")
    save_frame(storage, TABLE_DIR / "q3_storage_4h.csv")
    save_frame(table1, TABLE_DIR / "q3_specified_date_table1.csv")
    save_frame(table2, TABLE_DIR / "q3_specified_date_table2.csv")
    save_frame(table3, TABLE_DIR / "q3_specified_date_table3.csv")

    comparison = _comparison_frame(runs)
    save_frame(comparison, TABLE_DIR / "strategy_comparison.csv")
    save_frame(comparison[comparison["role"] == "diagnostic_frequency"], TABLE_DIR / "frequency_comparison.csv")
    save_frame(comparison[comparison["role"] == "diagnostic_mechanism"], TABLE_DIR / "mechanism_comparison.csv")
    if suite == "full":
        save_frame(comparison[comparison["role"] == "sensitivity"], TABLE_DIR / "sensitivity_comparison.csv")

    write_json(METRIC_DIR / "main_metrics.json", main_run.metrics)
    write_json(METRIC_DIR / "baseline_metrics.json", baseline_run.metrics)
    write_json(METRIC_DIR / "validation_checks.json", main_run.validation)
    write_json(METRIC_DIR / "forecast_diagnostics.json", forecast_diagnostics(data, official))

    method_entries = []
    for run in runs:
        method_entries.append(
            {
                "method_id": run.config.strategy_id,
                "role": run.config.role,
                "script": "code/Q3/q3_model.py",
                "status": "success",
                "execution_time_seconds": run.runtime_seconds,
                "input_files": [str(PRICE_PATH.relative_to(PROJECT_ROOT.parent)), str(ACTUAL_PATH.relative_to(PROJECT_ROOT.parent)), str(FORECAST_PATH.relative_to(PROJECT_ROOT.parent))],
                "output_files": [],
                "figure_files": [],
                "metrics_summary": run.metrics,
                "warnings": ["减购退款及逐次相对上一有效合同结算为Q3修订稿采用的解释性假设"],
                "errors": [],
            }
        )
    q_profiles = evaluation_ledger.pivot(index="date", columns="interval_index", values="original_contract_kwh")
    y_profiles = evaluation_ledger.pivot(index="date", columns="interval_index", values="final_contract_kwh")
    run_summary = {
        "schema_version": 1,
        "question": "Q3",
        "round": "round4",
        "implementation_target": "python",
        "random_seed": RANDOM_SEED,
        "approved_decision_id": "q3_method_choice_incremental_settlement_v4",
        "suite": suite,
        "methods": method_entries,
        "comparison": {
            "main_minus_baseline_cost_yuan": main_run.metrics["actual_total_cost_yuan"] - baseline_run.metrics["actual_total_cost_yuan"],
            "main_cost_change_vs_baseline_pct": (
                main_run.metrics["actual_total_cost_yuan"] / baseline_run.metrics["actual_total_cost_yuan"] - 1.0
            ) * 100.0,
        },
        "fallback_trigger": {
            "fallback_id": "F1",
            "condition": "any DA/TR stage failure or hard constraint violation",
            "observed": any(run.metrics["fallback_event_count"] for run in runs),
            "evidence": {run.config.strategy_id: run.validation["fallback_events"] for run in runs if run.validation["fallback_events"]},
        },
        "input_hashes": data.hashes,
        "environment": environment_info(),
        "validation": {run.config.strategy_id: run.validation for run in runs},
        "output_degeneracy": {
            "unique_original_contract_profiles_rounded_6dp": int(
                len(np.unique(np.round(q_profiles.to_numpy(), 6), axis=0))
            ),
            "unique_final_contract_profiles_rounded_6dp": int(
                len(np.unique(np.round(y_profiles.to_numpy(), 6), axis=0))
            ),
            "nonzero_adjustment_interval_count": int(
                np.count_nonzero(
                    np.abs(
                        evaluation_ledger["final_contract_kwh"].to_numpy()
                        - evaluation_ledger["original_contract_kwh"].to_numpy()
                    )
                    > 1e-6
                )
            ),
            "emergency_event_count": int(len(events)),
            "adjustment_release_interval_rows": int(len(evaluation_adjustments)),
            "adjustment_changed_rows": int(evaluation_adjustments["changed"].sum()),
        },
        "output_files": [
            "tables/q3_main_interval_ledger.csv",
            "tables/q3_adjustment_events.csv",
            "tables/q3_main_daily_summary.csv",
            "tables/q3_baseline_daily_summary.csv",
            "tables/q3_emergency_events.csv",
            "tables/q3_storage_4h.csv",
            "tables/q3_specified_date_table1.csv",
            "tables/q3_specified_date_table2.csv",
            "tables/q3_specified_date_table3.csv",
            "tables/strategy_comparison.csv",
            "metrics/main_metrics.json",
            "metrics/baseline_metrics.json",
            "metrics/validation_checks.json",
            "metrics/forecast_diagnostics.json",
        ],
    }
    write_json(ROUND_DIR / "run_summary.json", run_summary)


def main() -> None:
    parser = argparse.ArgumentParser(description="Solve Q3 rolling purchase and storage strategy")
    parser.add_argument("--suite", choices=("core", "sensitivity", "full"), default="core")
    args = parser.parse_args()
    run_suite(args.suite)


if __name__ == "__main__":
    raise SystemExit("请从辅助材料根目录运行 run.py，选择对应问题。")
